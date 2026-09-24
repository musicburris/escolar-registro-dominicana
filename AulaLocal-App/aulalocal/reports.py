"""Documentos y exportaciones generados localmente, sin plantillas remotas."""
import io, json, zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape
import reportlab
from PIL import Image as PILImage, ImageOps
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image as RLImage
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .catalog import MODULES
from .core import AppError,age_text
from .backup import atomic_write

FONT_DIR=Path(reportlab.__file__).parent/'fonts'
pdfmetrics.registerFont(TTFont('AulaSans',str(FONT_DIR/'Vera.ttf')))
pdfmetrics.registerFont(TTFont('AulaSans-Bold',str(FONT_DIR/'VeraBd.ttf')))
pdfmetrics.registerFontFamily('AulaSans',normal='AulaSans',bold='AulaSans-Bold',italic='AulaSans',boldItalic='AulaSans-Bold')
BLUE=colors.HexColor('#183849')
STYLES=getSampleStyleSheet()
for name in STYLES.byName:
    STYLES[name].fontName='AulaSans-Bold' if name.startswith('Heading') or name=='Title' else 'AulaSans'
STYLES.add(ParagraphStyle(name='CellAula',fontName='AulaSans',fontSize=8,leading=11,spaceAfter=0,wordWrap='CJK'))
STYLES.add(ParagraphStyle(name='LetterAula',fontName='AulaSans',fontSize=11,leading=18,spaceAfter=14))

def p(text,style='BodyText'): return Paragraph(escape(str(text)).replace('\n','<br/>'),STYLES[style])
def institution(store,year):
    y=store.year(year)
    frozen=json.loads(y.get('institution','{}'))
    return frozen or store.settings()

def title_block(store,year,title):
    s=institution(store,year); y=store.year(year)
    return [p(s.get('centro',''),'Title'),p(f"Código: {s.get('codigo','')} · Año escolar {y['label']} · {y['status']}",'Normal'),Spacer(1,16),p(title,'Heading1'),p(f'Fecha de emisión: {date.today().strftime("%d/%m/%Y")}','Normal'),Spacer(1,18)]
def footer(canvas,doc):
    canvas.saveState(); canvas.setStrokeColor(colors.HexColor('#d9e2e8')); canvas.line(36,34,doc.pagesize[0]-36,34)
    canvas.setFont('AulaSans',8); canvas.setFillColor(BLUE)
    canvas.drawString(36,22,'AulaLocal · Registro institucional'); canvas.drawRightString(doc.pagesize[0]-36,22,f'Página {doc.page}'); canvas.restoreState()
def pdf_bytes(story,wide=False):
    out=io.BytesIO(); doc=SimpleDocTemplate(out,pagesize=landscape(A4) if wide else A4,rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=48,title='AulaLocal - documento institucional')
    doc.build(story,onFirstPage=footer,onLaterPages=footer); return out.getvalue()
def table(headers,rows,width):
    values=[[p(h,'CellAula') for h in headers]]+[[p(v,'CellAula') for v in row] for row in rows]
    if not rows: values.append([p('Sin registros','CellAula')]+['']*(len(headers)-1))
    t=Table(values,colWidths=[width/len(headers)]*len(headers),repeatRows=1,hAlign='LEFT',splitInRow=1)
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dcebf0')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f3f6f8')]),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8),('LINEBELOW',(0,0),(-1,0),1,BLUE)]))
    return t

def readable_rows(store,kind,year,search='',filters=None):
    records=store.list_records(kind,year,search); result=[]
    filters={key:str(value) for key,value in (filters or {}).items() if value not in (None,'')}
    # Se consulta el mismo año; no resolver IDs en otros ciclos.
    refs={}
    if kind=='attendance': refs={r['id']:r['data']['nombre'] for r in store.list_records('staff',year)}
    if kind=='receipts': refs={r['id']:r['data']['nombre'] for r in store.list_records('menus',year,include_deleted=True)}
    for row in records:
        data=dict(row['data']); data['id']=row['id']
        if any(str(data.get(key,''))!=value for key,value in filters.items()):continue
        for f in MODULES[kind][1]:
            if f.type in ('staff_ref','menu_ref'): data[f.key]=refs.get(data[f.key],str(data[f.key]))
        if kind=='students':data['edad']=age_text(data['nacimiento'])
        result.append(data)
    return result

COLUMNS={'students':['codigo','nombre','nacimiento','edad','nivel','grado','seccion','estado'],'staff':['codigo','reloj_id','nombre','documento','cargo','ingreso','vinculo','telefono','estado'],
'attendance':['personal','fecha','estado','entrada','salida'],'menus':['nombre','ciclo','dia','servicio','descripcion'],
'receipts':['fecha','hora','servicio','suplidor','cantidad','rechazadas','incidencias','observaciones'],'expenses':['fecha','concepto','suplidor','factura','monto','estado'],
'facilities':['fecha','area','condicion','accion','responsable','estado'],'visits':['fecha','hora_entrada','hora_salida','nombre','institucion','motivo','recibido_por']}

FILTER_LABELS={'nivel':'Nivel','grado':'Grado','seccion':'Sección','estado':'Estado'}
def filter_description(filters=None):
    values=[f"{FILTER_LABELS.get(key,key.title())}: {value}" for key,value in (filters or {}).items() if value not in (None,'')]
    return ' · '.join(values)

def report_pdf(store,kind,year,path,search='',filters=None):
    rows=readable_rows(store,kind,year,search,filters); cols=COLUMNS[kind]; fields={f.key:f for f in MODULES[kind][1]};labels={k:f.label for k,f in fields.items()};labels['edad']='Edad actual';scope=filter_description(filters)
    story=title_block(store,year,MODULES[kind][0]); story.append(p(f'Registros incluidos: {len(rows)}'+(f' · Búsqueda: {search}' if search else '')+(f' · {scope}' if scope else ''))); story.append(Spacer(1,12))
    if kind=='expenses':
        total=sum((Decimal(r['monto']) for r in rows if r['estado']!='Anulado'),Decimal(0)); story.append(p(f'Total no anulado: RD$ {total:,.2f}','Heading2'))
    story.append(table([labels[c] for c in cols],[[r.get(c,'') for c in cols] for r in rows],landscape(A4)[0]-72))
    atomic_write(path,pdf_bytes(story,True)); store.log_export(kind,year)

def record_pdf(store,kind,year,rid,path):
    r=store.record(kind,year,rid); data=next(x for x in readable_rows(store,kind,year) if x['id']==rid)
    story=title_block(store,year,'Expediente / detalle del registro')
    story.append(p(f'{MODULES[kind][0]} · Registro {rid} · Versión {r["version"]}','Heading2'))
    for f in MODULES[kind][1]:
        story.extend([p(f.label,'Heading3'),p(data.get(f.key,'') or 'No registrado','LetterAula')])
    if kind=='students':
        profile=store.student_profile(year,r['data']);story.extend([p('Edad actual','Heading3'),p(profile['edad'],'LetterAula'),p('Edad al iniciar el año escolar','Heading3'),p(profile['edad_inicio_año'],'LetterAula'),p('Rango configurado para el grado','Heading3'),p(profile['rango_esperado'],'LetterAula')])
    at=store.attachments(kind,year,rid)
    if at:
        story.append(p('Adjuntos asociados','Heading2'))
        for a in at: story.append(p(a['name']))
    atomic_write(path,pdf_bytes(story)); store.log_export(kind,year)

def xlsx_bytes(store,kind,year,search='',filters=None):
    rows=readable_rows(store,kind,year,search,filters); fields=MODULES[kind][1]
    wb=Workbook(); ws=wb.active; ws.title=MODULES[kind][0][:31]
    scope=filter_description(filters)
    ws.append([institution(store,year)['centro']]); ws.append([f"Año escolar {store.year(year)['label']} · {MODULES[kind][0]}"+(f' · {scope}' if scope else '')]); ws.append([f.label for f in fields])
    for row in rows:
        values=[]
        for f in fields:
            v=row.get(f.key,'')
            if v!='' and f.type=='money': v=Decimal(v)
            elif v!='' and f.type=='date': v=date.fromisoformat(v)
            values.append(v)
        ws.append(values)
    # Las entradas de usuario se guardan como texto, nunca como fórmulas Excel.
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str): cell.data_type='s'
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    for cell in ws[3]: cell.fill=PatternFill('solid',fgColor='183849'); cell.font=Font(color='FFFFFF',bold=True)
    ws['A1'].font=Font(size=16,bold=True,color='183849'); ws['A2'].font=Font(size=11,color='506674')
    for i,f in enumerate(fields,1):
        ws.column_dimensions[get_column_letter(i)].width=34 if f.type=='long' or f.key in ('nombre','concepto') else 22
        for row in range(4,len(rows)+4):
            c=ws.cell(row,i)
            if f.type=='money': c.number_format='"RD$ "#,##0.00'
            elif f.type=='integer': c.number_format='#,##0'
            elif f.type=='date': c.number_format='dd/mm/yyyy'
    ws.freeze_panes='C4'; ws.auto_filter.ref=f'A3:{get_column_letter(len(fields))}{max(3,len(rows)+3)}'
    ws.sheet_view.showGridLines=False; ws.print_title_rows='1:3'; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A4; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    out=io.BytesIO(); wb.save(out); return out.getvalue()
def report_xlsx(store,kind,year,path,search='',filters=None):
    atomic_write(path,xlsx_bytes(store,kind,year,search,filters)); store.log_export(kind,year)

WEEKDAYS=('Lunes','Martes','Miércoles','Jueves','Viernes','Sábado','Domingo')
def _worked_minutes(entry,exit_time):
    if not entry or not exit_time:return 0
    start=int(entry[:2])*60+int(entry[3:]);end=int(exit_time[:2])*60+int(exit_time[3:])
    return max(0,end-start)

def attendance_report_data(store,year,start='',end='',search=''):
    y=store.year(year);start=start or y['start_date'];end=end or y['end_date']
    try:
        first=date.fromisoformat(start);last=date.fromisoformat(end)
        if first>last or start<y['start_date'] or end>y['end_date']:raise ValueError()
    except ValueError:raise AppError('El período del informe debe estar dentro del año escolar y usar AAAA-MM-DD.')
    staff={row['id']:row['data'] for row in store.list_records('staff',year)};detail=[]
    for row in store.list_records('attendance',year):
        data=row['data'];person=staff.get(data['personal'])
        if not person or not (start<=data['fecha']<=end):continue
        haystack=' '.join(map(str,[*data.values(),*person.values()])).casefold()
        if search and search.casefold() not in haystack:continue
        minutes=_worked_minutes(data.get('entrada',''),data.get('salida',''))
        detail.append({'codigo':person.get('codigo',''),'reloj_id':person.get('reloj_id',''),'nombre':person.get('nombre',''),'cargo':person.get('cargo',''),'fecha':data['fecha'],'dia':WEEKDAYS[date.fromisoformat(data['fecha']).weekday()],'entrada':data.get('entrada',''),'salida':data.get('salida',''),'minutos':minutes,'duracion':f'{minutes//60} h {minutes%60:02d} min' if minutes else '','estado':data.get('estado',''),'notas':data.get('notas','')})
    detail.sort(key=lambda item:(item['fecha'],item['nombre']));summary={}
    for item in detail:
        target=summary.setdefault(item['codigo'],{'codigo':item['codigo'],'reloj_id':item['reloj_id'],'nombre':item['nombre'],'cargo':item['cargo'],'dias':0,'presentes':0,'tardanzas':0,'ausencias':0,'permisos':0,'licencias':0,'minutos':0})
        target['dias']+=1;target['minutos']+=item['minutos'];target['presentes']+=item['estado']=='Presente';target['tardanzas']+=item['estado']=='Tardanza';target['ausencias']+=item['estado']=='Ausente';target['permisos']+=item['estado']=='Permiso';target['licencias']+=item['estado']=='Licencia'
    return {'start':start,'end':end,'detail':detail,'summary':sorted(summary.values(),key=lambda item:item['nombre'])}

def report_attendance_pdf(store,year,path,start='',end='',search=''):
    data=attendance_report_data(store,year,start,end,search);story=title_block(store,year,'Informe de asistencia del personal')
    story.extend([p(f"Período: {data['start']} al {data['end']} · Jornadas registradas: {len(data['detail'])}"),Spacer(1,12),p('Resumen por persona','Heading2')])
    summary_rows=[[r['codigo'],r['nombre'],r['cargo'],r['dias'],r['presentes'],r['tardanzas'],r['ausencias'],r['permisos'],f"{r['minutos']//60} h {r['minutos']%60:02d} min"] for r in data['summary']]
    story.append(table(['Código','Personal','Cargo','Días','Presente','Tardanza','Ausente','Permiso','Horas'],summary_rows,landscape(A4)[0]-72));story.extend([PageBreak(),p('Detalle diario','Heading2')])
    detail_rows=[[r['codigo'],r['nombre'],r['fecha'],r['dia'],r['entrada'],r['salida'],r['duracion'],r['estado'],r['notas']] for r in data['detail']]
    story.append(table(['Código','Personal','Fecha','Día','Entrada','Salida','Duración','Estado','Observaciones'],detail_rows,landscape(A4)[0]-72));atomic_write(path,pdf_bytes(story,True));store.log_export('attendance',year)

def attendance_xlsx_bytes(store,year,start='',end='',search=''):
    data=attendance_report_data(store,year,start,end,search);wb=Workbook();summary=wb.active;summary.title='Resumen';detail=wb.create_sheet('Detalle');center=institution(store,year)['centro'];period=f"{data['start']} al {data['end']}"
    summary.append([center]);summary.append([f'Informe de asistencia · {period}']);summary.append(['Código','ID ponchador','Personal','Cargo','Días registrados','Presente','Tardanza','Ausente','Permiso','Licencia','Horas trabajadas'])
    for row in data['summary']:summary.append([row['codigo'],row['reloj_id'],row['nombre'],row['cargo'],row['dias'],row['presentes'],row['tardanzas'],row['ausencias'],row['permisos'],row['licencias'],row['minutos']/60])
    detail.append([center]);detail.append([f'Detalle diario · {period}']);detail.append(['Código','ID ponchador','Personal','Cargo','Fecha','Día','Entrada','Salida','Duración','Estado','Observaciones'])
    for row in data['detail']:detail.append([row['codigo'],row['reloj_id'],row['nombre'],row['cargo'],date.fromisoformat(row['fecha']),row['dia'],row['entrada'],row['salida'],row['duracion'],row['estado'],row['notas']])
    for ws in (summary,detail):
        for cell in ws[3]:cell.fill=PatternFill('solid',fgColor='183849');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='top')
        ws['A1'].font=Font(size=16,bold=True,color='183849');ws['A2'].font=Font(color='506674');ws.freeze_panes='C4';ws.auto_filter.ref=f'A3:{get_column_letter(ws.max_column)}{max(3,ws.max_row)}';ws.sheet_view.showGridLines=False;ws.print_title_rows='1:3';ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True
        for column in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(column)].width=34 if column in (3,4,11) else 18
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top')
    for cell in summary['K'][3:]:cell.number_format='0.00'
    for cell in detail['E'][3:]:cell.number_format='dd/mm/yyyy'
    out=io.BytesIO();wb.save(out);return out.getvalue()

def report_attendance_xlsx(store,year,path,start='',end='',search=''):
    atomic_write(path,attendance_xlsx_bytes(store,year,start,end,search));store.log_export('attendance',year)

def food_report_data(store,year,start='',end='',service='',search=''):
    y=store.year(year);start=start or y['start_date'];end=end or y['end_date']
    try:
        first=date.fromisoformat(start);last=date.fromisoformat(end)
        if first>last or start<y['start_date'] or end>y['end_date']:raise ValueError()
    except ValueError:raise AppError('El período del informe debe estar dentro del año escolar y usar una fecha válida.')
    configured=[item['name'] for item in store.food_services(False)];used={row['data'].get('servicio','') for row in store.list_records('receipts',year)};services=configured+sorted(used-set(configured)-{''})
    if service and service not in services:raise AppError('Seleccione un tipo de alimentación válido.')
    menus={row['id']:row['data'] for row in store.list_records('menus',year,include_deleted=True)};detail=[];needle=search.casefold().strip()
    for row in store.list_records('receipts',year):
        data=row['data'];menu=menus.get(data['menu'],{})
        if not start<=data['fecha']<=end or service and data['servicio']!=service:continue
        haystack=' '.join(map(str,[*data.values(),menu.get('nombre',''),menu.get('descripcion','')])).casefold()
        if needle and needle not in haystack:continue
        evidence=[item for item in store.attachments('receipts',year,row['id']) if item['mime'].startswith('image/')];accepted=data['cantidad']-data['rechazadas'] if data['estado']!='Anulado' else 0
        detail.append({'id':row['id'],'fecha':data['fecha'],'dia':WEEKDAYS[date.fromisoformat(data['fecha']).weekday()],'hora':data['hora'],'servicio':data['servicio'],'menu':menu.get('nombre','Menú no disponible'),'descripcion':menu.get('descripcion',''),'suplidor':data['suplidor'],'cantidad':data['cantidad'],'rechazadas':data['rechazadas'],'aceptadas':accepted,'responsable':data['responsable'],'incidencias':data.get('incidencias',''),'observaciones':data.get('observaciones',''),'estado':data['estado'],'photos':evidence,'photo_count':len(evidence)})
    detail.sort(key=lambda item:(item['fecha'],item['hora'],item['servicio']));summary=[]
    for name in services:
        rows=[item for item in detail if item['servicio']==name]
        if rows or not service:summary.append({'servicio':name,'entregas':len(rows),'recibidas':sum(item['cantidad'] for item in rows if item['estado']!='Anulado'),'rechazadas':sum(item['rechazadas'] for item in rows if item['estado']!='Anulado'),'aceptadas':sum(item['aceptadas'] for item in rows),'fotos':sum(item['photo_count'] for item in rows)})
    return {'start':start,'end':end,'service':service or 'Todos','detail':detail,'summary':summary,'photo_count':sum(item['photo_count'] for item in detail)}

def _pdf_photo(content,name,max_width=340,max_height=230):
    try:
        source=PILImage.open(io.BytesIO(content));source=ImageOps.exif_transpose(source);source.thumbnail((1800,1400),PILImage.Resampling.LANCZOS)
        if source.mode not in ('RGB','L'):source=source.convert('RGB')
        converted=io.BytesIO();source.save(converted,'JPEG',quality=91,optimize=True);converted.seek(0);width,height=source.size;scale=min(max_width/width,max_height/height);image=RLImage(converted,width=width*scale,height=height*scale);image.hAlign='CENTER'
        return [image,p(name,'CellAula')]
    except Exception:return [p(f'{name}: la imagen no pudo incorporarse; el original permanece guardado.','CellAula')]

def _food_photo_grid(store,year,item,width):
    cells=[]
    for photo in item['photos']:
        content=store.attachment_bytes('receipts',year,item['id'],photo['id'],False);cells.append(_pdf_photo(content,photo['name'],width/2-20,230))
    if not cells:return p('Sin fotografías adjuntas para esta entrega.','Normal')
    rows=[]
    for index in range(0,len(cells),2):rows.append(cells[index:index+2]+([[]] if index+1==len(cells) else []))
    gallery=Table(rows,colWidths=[width/2,width/2],hAlign='LEFT',splitByRow=1);gallery.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BOX',(0,0),(-1,-1),0.5,colors.HexColor('#DDE7EC')),('INNERGRID',(0,0),(-1,-1),0.5,colors.HexColor('#DDE7EC')),('PADDING',(0,0),(-1,-1),8)]));return gallery

def report_food_pdf(store,year,path,start='',end='',service='',search='',include_photos=False,layout='summary'):
    if layout not in ('summary','delivery'):raise AppError('Seleccione un formato de informe válido.')
    data=food_report_data(store,year,start,end,service,search);width=landscape(A4)[0]-72;title='Informe institucional de alimentación escolar' if layout=='summary' else 'Fichas de recepción de alimentación escolar';story=title_block(store,year,title)
    story.extend([p(f"Período: {data['start']} al {data['end']} · Servicio: {data['service']} · Entregas: {len(data['detail'])} · Fotografías registradas: {data['photo_count']}"),Spacer(1,12)])
    summary_rows=[[row['servicio'],row['entregas'],row['recibidas'],row['rechazadas'],row['aceptadas'],row['fotos']] for row in data['summary']];story.extend([p('Resumen de cantidades','Heading2'),table(['Servicio','Entregas','Recibidas','Rechazadas','Aceptadas','Fotos'],summary_rows,width),Spacer(1,16)])
    if layout=='summary':
        rows=[[r['fecha'],r['dia'],r['hora'],r['servicio'],r['menu'],r['suplidor'],r['cantidad'],r['rechazadas'],r['aceptadas'],r['estado'],r['photo_count']] for r in data['detail']];story.extend([p('Detalle de recepciones','Heading2'),table(['Fecha','Día','Hora','Servicio','Menú','Suplidor','Recibidas','Rechazadas','Aceptadas','Estado','Fotos'],rows,width)])
        notes=[r for r in data['detail'] if r['incidencias'] or r['observaciones']]
        if notes:story.extend([Spacer(1,16),p('Incidencias y observaciones','Heading2'),table(['Fecha','Servicio','Menú','Incidencia o diferencia','Observaciones'],[[r['fecha'],r['servicio'],r['menu'],r['incidencias'] or '—',r['observaciones'] or '—'] for r in notes],width)])
        if include_photos and data['detail']:story.extend([PageBreak(),p('Evidencias fotográficas','Heading1')])
    elif data['detail']:story.append(PageBreak())
    if layout=='delivery' or include_photos:
        for index,item in enumerate(data['detail']):
            if index:story.append(PageBreak())
            story.extend([p(f"{item['servicio']} · {item['fecha']} · {item['hora']}",'Heading1'),table(['Día','Menú','Suplidor','Recibidas','Rechazadas','Aceptadas','Recibido por','Estado'],[[item['dia'],item['menu'],item['suplidor'],item['cantidad'],item['rechazadas'],item['aceptadas'],item['responsable'],item['estado']]],width),Spacer(1,10)])
            if item['descripcion']:story.extend([p('Menú registrado','Heading3'),p(item['descripcion']),Spacer(1,6)])
            story.extend([p('Incidencia, diferencia o producto incorrecto','Heading3'),p(item['incidencias'] or 'Sin incidencias registradas.'),p('Observaciones adicionales','Heading3'),p(item['observaciones'] or 'Sin observaciones adicionales.'),Spacer(1,10)])
            if include_photos:story.extend([p(f"Fotografías de la entrega ({item['photo_count']})",'Heading2'),_food_photo_grid(store,year,item,width)])
    if include_photos and not data['photo_count']:story.extend([Spacer(1,18),p('El período seleccionado no contiene fotografías adjuntas.','Heading2')])
    atomic_write(path,pdf_bytes(story,True));store.log_export('receipts',year)

def food_xlsx_bytes(store,year,start='',end='',service='',search=''):
    data=food_report_data(store,year,start,end,service,search);wb=Workbook();summary=wb.active;summary.title='Resumen';detail=wb.create_sheet('Recepciones');center=institution(store,year)['centro'];period=f"{data['start']} al {data['end']} · {data['service']}"
    summary.append([center]);summary.append([f'Informe de alimentación escolar · {period}']);summary.append(['Servicio','Entregas','Cantidad recibida','Unidades rechazadas','Cantidad aceptada','Fotografías'])
    for row in data['summary']:summary.append([row['servicio'],row['entregas'],row['recibidas'],row['rechazadas'],row['aceptadas'],row['fotos']])
    detail.append([center]);detail.append([f'Detalle de recepciones · {period}']);detail.append(['Fecha','Día','Hora','Servicio','Menú','Descripción','Suplidor','Recibidas','Rechazadas','Aceptadas','Recibido por','Incidencia o diferencia','Observaciones','Estado','Fotografías'])
    for row in data['detail']:detail.append([date.fromisoformat(row['fecha']),row['dia'],row['hora'],row['servicio'],row['menu'],row['descripcion'],row['suplidor'],row['cantidad'],row['rechazadas'],row['aceptadas'],row['responsable'],row['incidencias'],row['observaciones'],row['estado'],row['photo_count']])
    for ws in (summary,detail):
        for cell in ws[3]:cell.fill=PatternFill('solid',fgColor='183849');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='top')
        ws['A1'].font=Font(size=16,bold=True,color='183849');ws['A2'].font=Font(color='506674');ws.freeze_panes='C4';ws.auto_filter.ref=f'A3:{get_column_letter(ws.max_column)}{max(3,ws.max_row)}';ws.sheet_view.showGridLines=False;ws.print_title_rows='1:3';ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True
        for column in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(column)].width=36 if column in (5,6,7,12,13) else 18
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top')
    for cell in detail['A'][3:]:cell.number_format='dd/mm/yyyy'
    out=io.BytesIO();wb.save(out);return out.getvalue()

def report_food_xlsx(store,year,path,start='',end='',service='',search=''):
    atomic_write(path,food_xlsx_bytes(store,year,start,end,service,search));store.log_export('receipts',year)

DEFAULT_TEMPLATES={
 'staff': 'A quien pueda interesar:\n\nPor medio de la presente, {centro} hace constar que {nombre}, documento {documento}, figura en los registros del año escolar {año} con el cargo de {cargo}, fecha de ingreso {ingreso}, vínculo {vinculo} y estado {estado}.\n\nEsta constancia se expide a solicitud de la persona interesada, en fecha {fecha_emision}.',
 'students': 'A quien pueda interesar:\n\nLa Dirección de {centro} certifica que {nombre}, código {codigo_estudiante}, figura en la matrícula del año escolar {año}, en {grado}, sección {seccion}, con estado {estado}.\n\nEsta certificación se expide a solicitud de la persona interesada, en fecha {fecha_emision}.'}

def document_text(store,kind,year,rid,template=None):
    if kind not in DEFAULT_TEMPLATES: raise AppError('Seleccione un estudiante o una persona del personal.')
    d=store.record(kind,year,rid)['data']; s=institution(store,year)
    values={**d,**s,'nombre':d['nombre'],'año':store.year(year)['label'],'codigo_estudiante':d.get('codigo',''),'fecha_emision':date.today().strftime('%d/%m/%Y')}
    text=template or store.settings().get('template_'+kind,DEFAULT_TEMPLATES[kind])
    try: return text.format_map(values)
    except (KeyError,ValueError,AttributeError,IndexError) as e: raise AppError('La plantilla contiene un campo no válido. Use los campos de la plantilla original.') from e

def letter_pdf(store,kind,year,rid,path,body=None):
    store.record(kind,year,rid)
    if store.user['role'] not in ('admin','direccion','secretaria'): raise AppError('No tiene permiso para emitir certificaciones.')
    title='CONSTANCIA DE TRABAJO' if kind=='staff' else 'CERTIFICACIÓN DE MATRÍCULA'
    story=title_block(store,year,title)
    for paragraph in (body or document_text(store,kind,year,rid)).split('\n\n'): story.append(p(paragraph,'LetterAula'))
    story.extend([Spacer(1,50),p('________________________________','Normal'),p(institution(store,year).get('director',''),'Normal'),p('Dirección del centro educativo','Normal'),Spacer(1,16),p('Firma y sello institucional','Normal')])
    atomic_write(path,pdf_bytes(story)); store.log_export(kind,year)

def export_year(store,year,path):
    store.require(admin=True); y=store.year(year)
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('LEEME.txt','Archivo histórico consultable. No restaura la aplicación. Para restaurar, use un respaldo .aulabackup.\nLos adjuntos contienen datos personales; proteja este archivo.\n')
        z.writestr('año.json',json.dumps(y,ensure_ascii=False,indent=2))
        z.writestr('centro.json',json.dumps(institution(store,year),ensure_ascii=False,indent=2))
        z.writestr('estructura_academica.json',json.dumps(store.academic_summary(year),ensure_ascii=False,indent=2))
        z.writestr('auditoria.json',json.dumps(store.audit(year),ensure_ascii=False,indent=2))
        for kind in MODULES:
            rows=store.list_records(kind,year,include_deleted=True) if kind=='menus' else store.list_records(kind,year)
            z.writestr(f'{kind}.json',json.dumps(rows,ensure_ascii=False,indent=2)); z.writestr(f'{kind}.xlsx',attendance_xlsx_bytes(store,year) if kind=='attendance' else food_xlsx_bytes(store,year) if kind=='receipts' else xlsx_bytes(store,kind,year))
            for r in rows:
                for a in store.attachments(kind,year,r['id']):
                    name=Path(a['name']).name
                    z.writestr(f"adjuntos/{kind}/{r['id']}/{a['id']}-{name}",store.attachment_bytes(kind,year,r['id'],a['id']))
    atomic_write(path,out.getvalue())
    with store.db: store._audit('exportar_año',year)
