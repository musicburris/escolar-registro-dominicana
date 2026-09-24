"""Lectura local y tolerante de exportaciones de relojes ponchadores."""
import csv
import re
import unicodedata
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


class PunchFormatError(ValueError):
    pass


def norm(value):
    return re.sub(
        r"[^a-z0-9]+",
        " ",
        unicodedata.normalize("NFKD", str(value or ""))
        .encode("ascii", "ignore")
        .decode()
        .lower(),
    ).strip()


ALIASES = {
    "clock_id": {
        "id ponchador", "id reloj", "id usuario", "user id", "userid",
        "enroll number", "enrollment number", "employee id", "id empleado",
        "numero empleado", "numero de empleado", "ac no", "ac no.", "pin",
    },
    "code": {
        "codigo", "codigo personal", "codigo del personal", "matricula",
        "numero nomina", "numero de nomina", "employee code", "staff code",
    },
    "document": {"documento", "cedula", "cedula documento", "identificacion"},
    "name": {
        "nombre", "nombre completo", "nombres y apellidos", "empleado",
        "personal", "employee", "employee name", "name",
    },
    "date": {"fecha", "dia", "date", "punch date", "attendance date"},
    "time": {"hora", "marcacion", "marca", "punch", "punch time", "time"},
    "datetime": {
        "fecha hora", "fecha y hora", "fecha de marcacion", "hora de marcacion",
        "datetime", "date time", "punch datetime", "timestamp", "check time",
    },
    "entry": {
        "entrada", "hora entrada", "hora de entrada", "check in", "clock in",
        "in", "first punch", "primera marcacion",
    },
    "exit": {
        "salida", "hora salida", "hora de salida", "check out", "clock out",
        "out", "last punch", "ultima marcacion",
    },
    "direction": {
        "tipo", "tipo marcacion", "tipo de marcacion", "evento", "estado",
        "direccion", "direction", "punch state", "verify state",
    },
}

HEADER_MAP = {norm(alias): key for key, aliases in ALIASES.items() for alias in aliases}
ENTRY_WORDS = {"entrada", "in", "check in", "clock in", "ingreso", "entrada normal"}
EXIT_WORDS = {"salida", "out", "check out", "clock out", "egreso", "salida normal"}


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_rows(path):
    source = Path(path)
    if not source.is_file():
        raise PunchFormatError("No se encontró el archivo seleccionado.")
    if source.suffix.lower() == ".csv":
        for encoding in ("utf-8-sig", "latin-1"):
            try:
                with source.open(encoding=encoding, newline="") as handle:
                    sample = handle.read(8192)
                    handle.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                    except csv.Error:
                        dialect = csv.excel
                    return list(csv.reader(handle, dialect))
            except UnicodeDecodeError:
                continue
            except OSError as exc:
                raise PunchFormatError("No se pudo abrir el archivo CSV.") from exc
        raise PunchFormatError("No se pudo leer el texto del archivo CSV.")
    if source.suffix.lower() == ".xlsx":
        try:
            book = load_workbook(source, read_only=True, data_only=True)
            sheet = book.active
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
            raise PunchFormatError("El archivo Excel está dañado o no es un .xlsx válido.") from exc
    raise PunchFormatError("Seleccione un archivo Excel .xlsx o CSV.")


def _header_key(value):
    text = norm(value)
    if text in HEADER_MAP:
        return HEADER_MAP[text]
    matches = {key for alias, key in HEADER_MAP.items() if len(alias) >= 4 and alias in text}
    return next(iter(matches)) if len(matches) == 1 else None


def _header(raw):
    best = None
    for index, row in enumerate(raw[:25]):
        mapped = [_header_key(value) for value in row]
        found = {value for value in mapped if value}
        identity = found & {"clock_id", "code", "document", "name"}
        temporal = found & {"date", "time", "datetime", "entry", "exit"}
        score = len(found) + 2 * bool(identity) + 2 * bool(temporal)
        if identity and temporal and (best is None or score > best[0]):
            best = (score, index, mapped)
    if best is None:
        raise PunchFormatError(
            "No se reconocieron las columnas del ponchador. Incluya identificación o nombre, y fecha/hora, entrada o salida."
        )
    return best[1], best[2]


def _parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    text = _cell_text(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


def _parse_time(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, (int, float)):
        fraction = float(value) % 1
        minutes = round(fraction * 24 * 60) % (24 * 60)
        return time(minutes // 60, minutes % 60)
    text = _cell_text(value).upper().replace(".", "")
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M:%S %p", "%H%M"):
        try:
            return datetime.strptime(text, fmt).time().replace(second=0, microsecond=0)
        except ValueError:
            pass
    return None


def _parse_datetime(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, date):
        return datetime.combine(value, time())
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted if isinstance(converted, datetime) else datetime.combine(converted, time())
        except (TypeError, ValueError, OverflowError):
            return None
    text = _cell_text(value).upper().replace(".", "")
    for fmt in (
        "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M",
        "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d/%m/%Y %I:%M %p",
        "%m/%d/%Y %I:%M %p", "%m/%d/%Y %H:%M",
    ):
        try:
            return datetime.strptime(text, fmt).replace(second=0, microsecond=0)
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).replace(second=0, microsecond=0)
    except ValueError:
        return None


def _identity(data):
    values = {key: _cell_text(data.get(key)) for key in ("clock_id", "code", "document", "name")}
    for key in ("clock_id", "code", "document", "name"):
        if values[key]:
            return f"{key}:{norm(values[key])}", values
    return None, values


def parse_punch_file(path):
    """Devuelve jornadas agrupadas sin decidir qué persona de AulaLocal corresponde."""
    raw = _read_rows(path)
    if not raw:
        raise PunchFormatError("El archivo está vacío.")
    header_index, mapped = _header(raw)
    groups = {}
    errors = []
    source_marks = 0
    recognized_rows = 0

    def group_for(key, identities, day, row_number):
        group = groups.setdefault((key, day.isoformat()), {
            **identities, "date": day.isoformat(), "entries": [], "exits": [],
            "marks": [], "source_rows": [],
        })
        group["source_rows"].append(row_number)
        return group

    for row_number, row in enumerate(raw[header_index + 1:], header_index + 2):
        if not any(value not in (None, "") for value in row):
            continue
        data = {key: value for key, value in zip(mapped, row) if key and value not in (None, "")}
        key, identities = _identity(data)
        if not key:
            errors.append(f"Fila {row_number}: no tiene código, ID, documento ni nombre.")
            continue

        day = _parse_date(data.get("date"))
        entry_dt = _parse_datetime(data.get("entry"))
        exit_dt = _parse_datetime(data.get("exit"))
        entry = _parse_time(data.get("entry"))
        exit_time = _parse_time(data.get("exit"))
        stamp = _parse_datetime(data.get("datetime"))
        mark = _parse_time(data.get("time"))
        if not day:
            for candidate in (entry_dt, exit_dt, stamp):
                if candidate:
                    day = candidate.date()
                    break
        if not day:
            errors.append(f"Fila {row_number}: fecha no reconocida.")
            continue
        if entry_dt:
            entry = entry_dt.time()
        if exit_dt:
            exit_time = exit_dt.time()
        if stamp:
            mark = stamp.time()

        group = group_for(key, identities, day, row_number)
        row_marks = 0
        if entry:
            group["entries"].append(entry)
            row_marks += 1
        if exit_time:
            group["exits"].append(exit_time)
            row_marks += 1
        if mark:
            direction = norm(data.get("direction"))
            if direction in ENTRY_WORDS:
                group["entries"].append(mark)
            elif direction in EXIT_WORDS:
                group["exits"].append(mark)
            else:
                group["marks"].append(mark)
            row_marks += 1
        if not row_marks:
            errors.append(f"Fila {row_number}: hora no reconocida.")
            continue
        source_marks += row_marks
        recognized_rows += 1

    sessions = []
    for group in groups.values():
        all_marks = sorted(group["entries"] + group["exits"] + group["marks"])
        if not all_marks:
            continue
        entry = min(group["entries"]) if group["entries"] else all_marks[0]
        exit_time = max(group["exits"]) if group["exits"] else (all_marks[-1] if len(all_marks) > 1 else None)
        sessions.append({
            **{key: group[key] for key in ("clock_id", "code", "document", "name", "date", "source_rows")},
            "entry": entry.strftime("%H:%M") if entry else "",
            "exit": exit_time.strftime("%H:%M") if exit_time and exit_time != entry else "",
            "mark_count": len(all_marks),
        })
    if not sessions:
        detail = errors[0] if errors else "No hay marcaciones utilizables."
        raise PunchFormatError("No se reconocieron jornadas en el archivo. " + detail)
    sessions.sort(key=lambda item: (item["date"], norm(item["name"]), norm(item["code"])))
    return {
        "header_row": header_index + 1,
        "recognized_columns": [key for key in mapped if key],
        "source_rows": recognized_rows,
        "source_marks": source_marks,
        "sessions": sessions,
        "errors": errors,
    }
