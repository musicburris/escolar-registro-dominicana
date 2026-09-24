"""Capa de aplicación local. La UI no escribe SQL; todas las operaciones pasan aquí."""
import csv, hashlib, hmac, json, os, re, secrets, sqlite3, tempfile, time, unicodedata
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from .catalog import MODULES, ACCESS, ROLES
from .punch import PunchFormatError, parse_punch_file

class AppError(Exception): pass

def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')
def local_now(): return datetime.now().astimezone()
def timezone_label():
    current=local_now(); name=current.tzname() or 'Hora local'; offset=current.strftime('%z')
    utc=f'UTC{offset[:3]}:{offset[3:]}' if offset else 'Hora local'
    return utc if name.startswith(('+','-')) or name.upper() in ('UTC','GMT') else f'{name} · {utc}'
def age_parts(born,on=None):
    born=date.fromisoformat(str(born)); on=date.fromisoformat(str(on)) if on else date.today()
    if born>on: raise AppError('La fecha de nacimiento no puede ser futura.')
    years=on.year-born.year-((on.month,on.day)<(born.month,born.day))
    months=(on.year-born.year)*12+on.month-born.month-(on.day<born.day)-years*12
    return years,max(0,months)
def age_text(born,on=None):
    years,months=age_parts(born,on); return f'{years} años' + (f', {months} meses' if months else '')
def password_hash(password, salt):
    return hashlib.pbkdf2_hmac('sha256', password.encode(), salt, 600_000)
def check_password(password):
    if len(password)<12 or len(password)>256:
        raise AppError('La contraseña debe tener entre 12 y 256 caracteres.')
def app_directory():
    import sys
    base=Path.home()/('Library/Application Support' if sys.platform=='darwin' else '.local/share')
    return base/'AulaLocal'

class Store:
    def __init__(self, root=None):
        self.root=Path(root or app_directory()); self.root.mkdir(parents=True,exist_ok=True)
        os.chmod(self.root,0o700)
        self.path=self.root/'aulalocal.sqlite3'
        self.db=sqlite3.connect(self.path,timeout=10)
        self.db.row_factory=sqlite3.Row
        self.db.execute('PRAGMA foreign_keys=ON')
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.execute('PRAGMA synchronous=FULL')
        self.db.execute('PRAGMA busy_timeout=10000')
        schema=Path(__file__).with_name('schema.sql').read_text()
        # Migrar sin perder datos y no abrir silenciosamente una base futura.
        if self.db.execute("SELECT 1 FROM sqlite_master WHERE name='meta'").fetchone():
            v=self.db.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()
            if not v or v[0] not in ('1','2'): raise AppError('Versión de datos no compatible.')
            if v[0]=='1': self._migrate_v1()
        self.db.executescript(schema); self.db.commit(); os.chmod(self.path,0o600)
        self._seed_academic_catalog()
        self._seed_food_services()
        self._seed_attendance_settings()
        self._normalize_legacy_students()
        self.user=None; self.last_active=0; self.session_hash=None
        self.resume_session()
    def _migrate_v1(self):
        self.db.execute('PRAGMA foreign_keys=OFF')
        self.db.executescript('''
        BEGIN IMMEDIATE;
        DROP TRIGGER IF EXISTS records_insert_lock; DROP TRIGGER IF EXISTS records_update_lock;
        DROP TRIGGER IF EXISTS records_identity; DROP TRIGGER IF EXISTS records_delete_lock;
        DROP TRIGGER IF EXISTS attachments_insert_lock; DROP TRIGGER IF EXISTS attachments_update_lock; DROP TRIGGER IF EXISTS attachments_delete_lock;
        ALTER TABLE attachments RENAME TO attachments_v1;
        ALTER TABLE records RENAME TO records_v1;
        CREATE TABLE records(id INTEGER PRIMARY KEY,year_id INTEGER NOT NULL REFERENCES years(id),kind TEXT NOT NULL,data TEXT NOT NULL CHECK(json_valid(data)),created_at TEXT NOT NULL,updated_at TEXT NOT NULL,version INTEGER NOT NULL DEFAULT 1,UNIQUE(id,year_id));
        INSERT INTO records SELECT * FROM records_v1;
        CREATE TABLE attachments(id INTEGER PRIMARY KEY,record_id INTEGER NOT NULL,year_id INTEGER NOT NULL,name TEXT NOT NULL,mime TEXT NOT NULL,sha256 TEXT NOT NULL,content BLOB NOT NULL,FOREIGN KEY(record_id,year_id) REFERENCES records(id,year_id));
        INSERT INTO attachments SELECT * FROM attachments_v1;
        DROP TABLE attachments_v1; DROP TABLE records_v1;
        UPDATE meta SET value='2' WHERE key='schema_version';
        COMMIT;''')
        self.db.execute('PRAGMA foreign_keys=ON')
    def _seed_academic_catalog(self):
        levels=[('Nivel Inicial',10),('Nivel Primario',20),('Nivel Secundario',30)]
        grades={'Nivel Inicial':[('Kínder',4,4),('Preprimario',5,5)],
          'Nivel Primario':[(f'{n}.º de Primaria',5+n,5+n) for n in range(1,7)],
          'Nivel Secundario':[(f'{n}.º de Secundaria',11+n,11+n) for n in range(1,7)]}
        with self.db:
            self.db.executemany('INSERT OR IGNORE INTO academic_levels(name,sort_order) VALUES(?,?)',levels)
            for level,items in grades.items():
                lid=self.db.execute('SELECT id FROM academic_levels WHERE name=?',(level,)).fetchone()[0]
                for order,(name,minimum,maximum) in enumerate(items,1):
                    self.db.execute('INSERT OR IGNORE INTO grades(level_id,name,min_age,max_age,sort_order) VALUES(?,?,?,?,?)',(lid,name,minimum,maximum,order))
            self.db.execute("INSERT OR IGNORE INTO sections(year_id,grade_id,name) SELECT y.id,g.id,'A' FROM years y CROSS JOIN grades g")
    def _seed_attendance_settings(self):
        with self.db:
            self.db.executemany('INSERT OR IGNORE INTO settings VALUES(?,?)',[
                ('attendance_start','08:00'),('attendance_end','16:00'),('attendance_grace','10')])
    def _seed_food_services(self):
        with self.db:
            self.db.executemany('INSERT OR IGNORE INTO food_services(name,created_at) VALUES(?,?)',
                [(name,now()) for name in ('Leche','Pan','Almuerzo')])
    @staticmethod
    def _norm(value):
        return re.sub(r'[^a-z0-9]+',' ',unicodedata.normalize('NFKD',str(value)).encode('ascii','ignore').decode().lower()).strip()
    def _match_level(self,value):
        needle=self._norm(value)
        aliases={'inicial':'nivel inicial','preescolar':'nivel inicial','primaria':'nivel primario',
                 'primario':'nivel primario','secundaria':'nivel secundario','secundario':'nivel secundario'}
        needle=aliases.get(needle,needle)
        levels=self.db.execute('SELECT * FROM academic_levels').fetchall()
        matches=[dict(level) for level in levels if self._norm(level['name'])==needle]
        return matches[0] if len(matches)==1 else None
    def _match_grade(self,value,level=None):
        needle=self._norm(value)
        aliases={'kinder':'Kínder','pre primaria':'Preprimario','preprimaria':'Preprimario','pre primario':'Preprimario','preprimario':'Preprimario'}
        if needle in aliases:needle=self._norm(aliases[needle])
        grades=self.db.execute('SELECT g.name,l.name level,g.id,g.min_age,g.max_age FROM grades g JOIN academic_levels l ON l.id=g.level_id').fetchall()
        matched_level=self._match_level(level) if level else None
        if matched_level:grades=[g for g in grades if g['level']==matched_level['name']]
        exact=[g for g in grades if self._norm(g['name'])==needle]
        if exact:return dict(exact[0])
        ordinals={'primero':'1','primer':'1','1ro':'1','segundo':'2','2do':'2','tercero':'3','3ro':'3',
                  'cuarto':'4','4to':'4','quinto':'5','5to':'5','sexto':'6','6to':'6'}
        words=needle.split()
        if words and words[0] in ordinals:words[0]=ordinals[words[0]];needle=' '.join(words)
        filler={'o','de','del','grado','nivel'}
        tokens={token for token in needle.split() if token not in filler}
        candidates=[g for g in grades if tokens and tokens<={token for token in self._norm(g['name']).split() if token not in filler}]
        return dict(candidates[0]) if len(candidates)==1 else None
    def _normalize_legacy_students(self):
        rows=self.db.execute("SELECT id,year_id,data FROM records WHERE kind='students'").fetchall()
        with self.db:
            for row in rows:
                data=json.loads(row['data'])
                grade=self._match_grade(data.get('grado',''))
                if not grade:continue
                data['nivel']=grade['level'];data['grado']=grade['name'];data['seccion']=str(data.get('seccion') or 'A').upper()
                self.db.execute('INSERT OR IGNORE INTO sections(year_id,grade_id,name) VALUES(?,?,?)',(row['year_id'],grade['id'],data['seccion']))
                self.db.execute('UPDATE records SET data=? WHERE id=?',(json.dumps(data,ensure_ascii=False),row['id']))
    def close(self): self.db.close()
    def initialized(self): return bool(self.db.execute('SELECT 1 FROM users LIMIT 1').fetchone())
    def _audit(self,action,year=None,detail=''):
        self.db.execute('INSERT INTO audit(at,user_id,action,year_id,detail) VALUES(?,?,?,?,?)',
            (now(),self.user['id'] if self.user else None,action,year,str(detail)))
    def bootstrap(self,username,password,name,school):
        if self.initialized(): raise AppError('La aplicación ya está configurada.')
        check_password(password)
        if not re.fullmatch(r'[a-zA-Z0-9_.-]{3,64}',username.strip()): raise AppError('El usuario debe tener de 3 a 64 letras, números, puntos o guiones.')
        if not all(str(x).strip() for x in [username,name,school]): raise AppError('Complete todos los datos.')
        salt=os.urandom(16)
        with self.db:
            self.db.execute('INSERT INTO users(username,name,role,salt,password_hash) VALUES(?,?,?,?,?)',
                (username.strip().lower(),name.strip(),'admin',salt,password_hash(password,salt)))
            self.db.executemany('INSERT INTO settings VALUES(?,?)',[('centro',school.strip()),('director',name.strip()),('codigo',''),('direccion',''),('telefono','')])
            self._audit('configuracion_inicial')
        self.login(username,password,remember=True)
        return self.create_recovery_key(password)
    def login(self,username,password,remember=True):
        if self.session_hash:
            with self.db:self.db.execute('DELETE FROM session_tokens WHERE token_hash=?',(self.session_hash,))
            self.session_hash=None;self.session_file.unlink(missing_ok=True)
        self.user=None
        row=self.db.execute('SELECT * FROM users WHERE username=?',(username.strip().lower(),)).fetchone()
        if row and row['locked_until']>time.time(): raise AppError('Cuenta bloqueada temporalmente. Espere 5 minutos.')
        salt=row['salt'] if row else b'0'*16
        valid=hmac.compare_digest(password_hash(password,salt),row['password_hash'] if row else b'0'*32)
        if not row or not valid or not row['active']:
            with self.db:
                if row:
                    fails=row['failures']+1
                    self.db.execute('UPDATE users SET failures=?,locked_until=? WHERE id=?',(fails,time.time()+300 if fails>=5 else 0,row['id']))
                self._audit('acceso_fallido')
            raise AppError('Usuario o contraseña incorrectos.')
        self.user={'id':row['id'],'name':row['name'],'role':row['role'],'username':row['username']}; self.last_active=time.monotonic()
        with self.db:
            self.db.execute('UPDATE users SET failures=0,locked_until=0 WHERE id=?',(row['id'],)); self._audit('inicio_sesion')
        if remember:self._create_session_token(row['id'])
        return self.user
    def logout(self):
        if self.user:
            with self.db:
                self._audit('cierre_sesion')
                if self.session_hash:self.db.execute('DELETE FROM session_tokens WHERE token_hash=?',(self.session_hash,))
        self.user=None; self.session_hash=None
        try:self.session_file.unlink(missing_ok=True)
        except OSError:pass
    @property
    def session_file(self): return self.root/'session.key'
    def _create_session_token(self,user_id):
        token=secrets.token_urlsafe(48); digest=hashlib.sha256(token.encode()).digest()
        with self.db:
            self.db.execute('DELETE FROM session_tokens WHERE user_id=?',(user_id,))
            self.db.execute('INSERT INTO session_tokens(user_id,token_hash,created_at,last_used) VALUES(?,?,?,?)',(user_id,digest,now(),now()))
        self.session_file.write_text(token,encoding='utf-8'); os.chmod(self.session_file,0o600); self.session_hash=digest
    def resume_session(self):
        try: token=self.session_file.read_text(encoding='utf-8').strip()
        except OSError:return False
        digest=hashlib.sha256(token.encode()).digest()
        row=self.db.execute('SELECT u.* FROM session_tokens s JOIN users u ON u.id=s.user_id WHERE s.token_hash=? AND u.active=1',(digest,)).fetchone()
        if not row:
            self.session_file.unlink(missing_ok=True); return False
        self.user={'id':row['id'],'name':row['name'],'role':row['role'],'username':row['username']}; self.last_active=time.monotonic(); self.session_hash=digest
        with self.db:self.db.execute('UPDATE session_tokens SET last_used=? WHERE token_hash=?',(now(),digest))
        return True
    def require(self,kind=None,write=False,admin=False):
        if not self.user: raise AppError('Inicie sesión para continuar.')
        current=self.db.execute('SELECT role,active FROM users WHERE id=?',(self.user['id'],)).fetchone()
        if not current or not current['active']:
            self.user=None; raise AppError('La cuenta está desactivada.')
        self.user['role']=current['role']; role=current['role']
        if admin and role!='admin': raise AppError('Esta acción requiere una cuenta administradora.')
        if kind and kind not in ACCESS[role]: raise AppError('No tiene acceso a este módulo.')
        if write and role=='consulta': raise AppError('Esta cuenta solo permite consultar.')
        self.last_active=time.monotonic()
    def verify_password(self,password):
        self.require()
        u=self.db.execute('SELECT * FROM users WHERE id=?',(self.user['id'],)).fetchone()
        if not hmac.compare_digest(password_hash(password,u['salt']),u['password_hash']): raise AppError('Contraseña incorrecta.')
    def create_recovery_key(self,password):
        self.verify_password(password); code='AULA-'+secrets.token_hex(4).upper()+'-'+secrets.token_hex(4).upper()
        salt=os.urandom(16)
        with self.db:
            self.db.execute('INSERT OR REPLACE INTO recovery_keys(user_id,salt,key_hash,created_at) VALUES(?,?,?,?)',(self.user['id'],salt,password_hash(code,salt),now()))
            self._audit('crear_clave_recuperacion',detail=f"usuario={self.user['username']}")
        return code
    def recover_access(self,code,new_username,new_password):
        check_password(new_password)
        if not re.fullmatch(r'[a-zA-Z0-9_.-]{3,64}',new_username.strip()): raise AppError('El nuevo usuario no es válido.')
        found=None
        for row in self.db.execute('SELECT * FROM recovery_keys'):
            if hmac.compare_digest(password_hash(code.strip().upper(),row['salt']),row['key_hash']): found=row; break
        if not found: raise AppError('El código de recuperación no es válido.')
        salt=os.urandom(16);new_code='AULA-'+secrets.token_hex(4).upper()+'-'+secrets.token_hex(4).upper();recovery_salt=os.urandom(16)
        try:
            with self.db:
                self.db.execute('UPDATE users SET username=?,salt=?,password_hash=?,failures=0,locked_until=0 WHERE id=?',(new_username.strip().lower(),salt,password_hash(new_password,salt),found['user_id']))
                self.db.execute('DELETE FROM session_tokens WHERE user_id=?',(found['user_id'],))
                self.db.execute('INSERT OR REPLACE INTO recovery_keys(user_id,salt,key_hash,created_at) VALUES(?,?,?,?)',(found['user_id'],recovery_salt,password_hash(new_code,recovery_salt),now()))
                self._audit('recuperar_acceso',detail=f"usuario_id={found['user_id']}")
        except sqlite3.IntegrityError as e: raise AppError('Ese nombre de usuario ya existe.') from e
        return new_code
    def reset_user_password(self,user_id,new_password):
        self.require(admin=True); check_password(new_password)
        if user_id==self.user['id']: raise AppError('Para tu propia cuenta usa Cambiar contraseña o el código de recuperación.')
        salt=os.urandom(16)
        with self.db:
            self.db.execute('UPDATE users SET salt=?,password_hash=?,failures=0,locked_until=0 WHERE id=?',(salt,password_hash(new_password,salt),user_id))
            self.db.execute('DELETE FROM session_tokens WHERE user_id=?',(user_id,)); self._audit('restablecer_contraseña',detail=f'usuario_id={user_id}')
    def users(self):
        self.require(admin=True)
        return [dict(x) for x in self.db.execute('SELECT id,username,name,role,active FROM users ORDER BY name')]
    def save_user(self,username,name,role,password='',active=True,user_id=None):
        self.require(admin=True)
        if role not in ROLES or not name.strip() or not re.fullmatch(r'[a-zA-Z0-9_.-]{3,64}',username): raise AppError('Datos de usuario no válidos. Usuario: 3–64 letras, números, punto o guion.')
        if user_id==self.user['id'] and (role!='admin' or not active): raise AppError('No puede desactivar o degradar su propia cuenta.')
        if password: check_password(password)
        if not user_id and not password: raise AppError('Defina una contraseña.')
        try:
            with self.db:
                if user_id:
                    if not self.db.execute('SELECT 1 FROM users WHERE id=?',(user_id,)).fetchone(): raise AppError('Usuario inexistente.')
                    self.db.execute('UPDATE users SET username=?,name=?,role=?,active=? WHERE id=?',(username.lower(),name.strip(),role,int(active),user_id))
                else:
                    salt=os.urandom(16)
                    user_id=self.db.execute('INSERT INTO users(username,name,role,salt,password_hash) VALUES(?,?,?,?,?)',(username.lower(),name.strip(),role,salt,password_hash(password,salt))).lastrowid
                if password:
                    salt=os.urandom(16)
                    self.db.execute('UPDATE users SET salt=?,password_hash=?,failures=0,locked_until=0 WHERE id=?',(salt,password_hash(password,salt),user_id))
                self._audit('guardar_usuario',detail=f'id={user_id}; rol={role}; activo={active}')
        except sqlite3.IntegrityError as e: raise AppError('El nombre de usuario ya existe.') from e
    def settings(self):
        self.require(); return dict(self.db.execute('SELECT key,value FROM settings'))
    def save_settings(self,values):
        self.require(admin=True)
        if not values.get('centro','').strip() or not values.get('director','').strip(): raise AppError('Centro y director son obligatorios.')
        with self.db:
            for key,value in values.items():
                if key not in {'centro','director','codigo','direccion','telefono','template_staff','template_students','attendance_start','attendance_end','attendance_grace'}: raise AppError('Configuración desconocida.')
                if key in ('attendance_start','attendance_end') and not re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d',str(value).strip()):raise AppError('El horario institucional debe usar HH:MM.')
                if key=='attendance_grace':
                    try:
                        grace=int(str(value).strip())
                        if not 0<=grace<=120:raise ValueError()
                    except ValueError:raise AppError('La tolerancia debe ser de 0 a 120 minutos.')
                self.db.execute('INSERT OR REPLACE INTO settings VALUES(?,?)',(key,str(value).strip()[:10000]))
            self._audit('actualizar_configuracion')
    def years(self):
        self.require(); return [dict(x) for x in self.db.execute('SELECT * FROM years ORDER BY label DESC')]
    def year(self,year):
        self.require(); row=self.db.execute('SELECT * FROM years WHERE id=?',(year,)).fetchone()
        if not row: raise AppError('Seleccione un año escolar válido.')
        return dict(row)
    def create_year(self,label,start,end):
        self.require(admin=True)
        match=re.fullmatch(r'(20\d{2})-(20\d{2})',label)
        if not match or int(match[2])!=int(match[1])+1: raise AppError('Use un año como 2026-2027.')
        try:
            a,b=date.fromisoformat(start),date.fromisoformat(end)
            if a>=b or a.year!=int(match[1]) or b.year!=int(match[2]): raise ValueError()
        except ValueError: raise AppError('Las fechas deben corresponder al año escolar indicado.')
        try:
            with self.db:
                y=self.db.execute('INSERT INTO years(label,start_date,end_date) VALUES(?,?,?)',(label,start,end)).lastrowid
                self.db.execute("INSERT OR IGNORE INTO sections(year_id,grade_id,name,capacity) SELECT ?,id,'A',NULL FROM grades WHERE active=1",(y,))
                self._audit('crear_año',y,label)
            return y
        except sqlite3.IntegrityError as e: raise AppError('Ese año ya existe.') from e
    def transition(self,year,status,reason,password):
        self.require(admin=True); self.verify_password(password); old=self.year(year)
        allowed={'abierto':{'congelado'},'congelado':{'abierto','archivado'},'archivado':{'congelado'}}
        if status not in allowed[old['status']]: raise AppError('Transición no permitida; el archivo histórico se descongela en dos etapas.')
        if len(reason.strip())<10: raise AppError('Escriba un motivo de al menos 10 caracteres.')
        with self.db:
            self.db.execute('UPDATE years SET status=?,institution=? WHERE id=?',(status,json.dumps(self.settings(),ensure_ascii=False) if old['status']=='abierto' else old['institution'],year)); self._audit('estado_año',year,f"{old['status']} → {status}: {reason.strip()}")
    def _open(self,year):
        if self.year(year)['status']!='abierto': raise AppError('El año está congelado o archivado; solo permite consultas.')
    def academic_levels(self,active_only=True):
        self.require(); sql='SELECT * FROM academic_levels'+(' WHERE active=1' if active_only else '')+' ORDER BY sort_order,name'
        return [dict(r) for r in self.db.execute(sql)]
    def grades(self,level_id=None,active_only=True):
        self.require(); where=[];args=[]
        if level_id is not None:where.append('g.level_id=?');args.append(level_id)
        if active_only:where.append('g.active=1')
        sql='SELECT g.*,l.name level FROM grades g JOIN academic_levels l ON l.id=g.level_id'+((' WHERE '+' AND '.join(where)) if where else '')+' ORDER BY l.sort_order,g.sort_order,g.name'
        return [dict(r) for r in self.db.execute(sql,args)]
    def sections(self,year,grade_id=None,active_only=True):
        self.year(year);where=['s.year_id=?'];args=[year]
        if grade_id is not None:where.append('s.grade_id=?');args.append(grade_id)
        if active_only:where.append('s.active=1')
        sql='SELECT s.*,g.name grade,l.name level FROM sections s JOIN grades g ON g.id=s.grade_id JOIN academic_levels l ON l.id=g.level_id WHERE '+' AND '.join(where)+' ORDER BY l.sort_order,g.sort_order,s.name'
        return [dict(r) for r in self.db.execute(sql,args)]
    def food_services(self,active_only=True):
        self.require('menus')
        sql='SELECT * FROM food_services'+(' WHERE active=1' if active_only else '')+' ORDER BY name COLLATE NOCASE'
        return [dict(row) for row in self.db.execute(sql)]
    def save_food_service(self,name):
        self.require('menus',write=True);name=' '.join(str(name).split())
        if not 2<=len(name)<=60:raise AppError('El tipo de alimentación debe tener entre 2 y 60 caracteres.')
        if not re.fullmatch(r"[\wÁÉÍÓÚÜÑáéíóúüñ .()'/-]+",name,re.UNICODE):raise AppError('El nombre contiene caracteres no permitidos.')
        existing=self.db.execute('SELECT id FROM food_services WHERE name=? COLLATE NOCASE',(name,)).fetchone()
        with self.db:
            if existing:
                service_id=existing['id'];self.db.execute('UPDATE food_services SET name=?,active=1 WHERE id=?',(name,service_id))
            else:service_id=self.db.execute('INSERT INTO food_services(name,created_at) VALUES(?,?)',(name,now())).lastrowid
            self._audit('guardar_tipo_alimentacion',detail=f'id={service_id}; nombre={name}')
        return service_id
    def set_food_service_active(self,service_id,active):
        self.require('menus',write=True)
        row=self.db.execute('SELECT * FROM food_services WHERE id=?',(service_id,)).fetchone()
        if not row:raise AppError('El tipo de alimentación no existe.')
        if not active and self.db.execute('SELECT count(*) FROM food_services WHERE active=1').fetchone()[0]<=1:raise AppError('Debe conservar al menos un tipo de alimentación activo.')
        with self.db:
            self.db.execute('UPDATE food_services SET active=? WHERE id=?',(int(bool(active)),service_id))
            self._audit('activar_tipo_alimentacion' if active else 'ocultar_tipo_alimentacion',detail=f"id={service_id}; nombre={row['name']}")
    def save_level(self,name,sort_order=0,level_id=None):
        self.require(admin=True);name=name.strip()
        if not name:raise AppError('Escriba el nombre del nivel.')
        try:
            with self.db:
                if level_id:self.db.execute('UPDATE academic_levels SET name=?,sort_order=? WHERE id=?',(name,int(sort_order),level_id))
                else:level_id=self.db.execute('INSERT INTO academic_levels(name,sort_order) VALUES(?,?)',(name,int(sort_order))).lastrowid
                self._audit('guardar_nivel',detail=f'id={level_id}; nombre={name}')
            return level_id
        except sqlite3.IntegrityError as e:raise AppError('Ese nivel ya existe.') from e
    def save_grade(self,level_id,name,min_age=None,max_age=None,sort_order=0,grade_id=None):
        self.require(admin=True);name=name.strip()
        if not self.db.execute('SELECT 1 FROM academic_levels WHERE id=?',(level_id,)).fetchone():raise AppError('Seleccione un nivel válido.')
        try:
            minimum=None if min_age in (None,'') else int(min_age);maximum=None if max_age in (None,'') else int(max_age)
            if minimum is not None and not 0<=minimum<=30:raise ValueError()
            if maximum is not None and not 0<=maximum<=30:raise ValueError()
            if minimum is not None and maximum is not None and minimum>maximum:raise ValueError()
        except ValueError:raise AppError('El rango de edad debe ser válido, entre 0 y 30 años.')
        try:
            with self.db:
                if grade_id:self.db.execute('UPDATE grades SET level_id=?,name=?,min_age=?,max_age=?,sort_order=? WHERE id=?',(level_id,name,minimum,maximum,int(sort_order),grade_id))
                else:grade_id=self.db.execute('INSERT INTO grades(level_id,name,min_age,max_age,sort_order) VALUES(?,?,?,?,?)',(level_id,name,minimum,maximum,int(sort_order))).lastrowid
                self._audit('guardar_grado',detail=f'id={grade_id}; nombre={name}')
            return grade_id
        except sqlite3.IntegrityError as e:raise AppError('Ese grado ya existe dentro del nivel.') from e
    def save_section(self,year,grade_id,name,capacity=None,section_id=None):
        self.require(admin=True);self._open(year);name=name.strip().upper()
        if not name:raise AppError('Escriba el nombre de la sección.')
        try:
            capacity=None if capacity in (None,'') else int(capacity)
            if capacity is not None and not 1<=capacity<=200:raise ValueError()
        except ValueError:raise AppError('La capacidad debe estar entre 1 y 200 estudiantes.')
        try:
            with self.db:
                if section_id:self.db.execute('UPDATE sections SET grade_id=?,name=?,capacity=? WHERE id=? AND year_id=?',(grade_id,name,capacity,section_id,year))
                else:section_id=self.db.execute('INSERT INTO sections(year_id,grade_id,name,capacity) VALUES(?,?,?,?)',(year,grade_id,name,capacity)).lastrowid
                self._audit('guardar_seccion',year,f'id={section_id}; grado={grade_id}; nombre={name}')
            return section_id
        except sqlite3.IntegrityError as e:raise AppError('Esa sección ya existe para el grado seleccionado.') from e
    def academic_summary(self,year):
        rows=[]
        for section in self.sections(year,active_only=False):
            count=self.db.execute("SELECT count(*) FROM records WHERE year_id=? AND kind='students' AND json_extract(data,'$.nivel')=? AND json_extract(data,'$.grado')=? AND json_extract(data,'$.seccion')=? AND json_extract(data,'$.estado')='Activo'",(year,section['level'],section['grade'],section['name'])).fetchone()[0]
            rows.append({**section,'students':count})
        return rows
    def list_records(self,kind,year,search='',include_deleted=False):
        self.require(kind); self.year(year)
        rows=[dict(x) for x in self.db.execute('SELECT * FROM records WHERE kind=? AND year_id=? ORDER BY id DESC',(kind,year))]
        for r in rows: r['data']=json.loads(r['data'])
        if kind=='menus' and not include_deleted:rows=[r for r in rows if not r['data'].get('_deleted')]
        if search: rows=[r for r in rows if search.casefold() in ' '.join(map(str,r['data'].values())).casefold()]
        return rows
    def record(self,kind,year,rid):
        self.require(kind); self.year(year)
        r=self.db.execute('SELECT * FROM records WHERE id=? AND year_id=? AND kind=?',(rid,year,kind)).fetchone()
        if not r: raise AppError('No se encontró el registro en este año.')
        out=dict(r); out['data']=json.loads(out['data']); return out
    def _validate(self,kind,year,data):
        if kind not in MODULES: raise AppError('Módulo desconocido.')
        y=self.year(year); out={}
        known={f.key for f in MODULES[kind][1]}
        if set(data)-known: raise AppError('El registro contiene campos desconocidos.')
        for f in MODULES[kind][1]:
            v=str(data.get(f.key,'')).strip()
            if f.required and not v: raise AppError(f'Complete: {f.label}.')
            if len(v)>10000: raise AppError(f'{f.label}: texto demasiado largo.')
            if not v: out[f.key]=''; continue
            if f.type=='choice' and v not in f.options: raise AppError(f'{f.label}: opción no válida.')
            if f.type=='food_service_ref' and not self.db.execute('SELECT 1 FROM food_services WHERE name=? COLLATE NOCASE',(v,)).fetchone():raise AppError('Seleccione un tipo de alimentación válido.')
            if f.type=='level_ref':
                if not self.db.execute('SELECT 1 FROM academic_levels WHERE name=? AND active=1',(v,)).fetchone():raise AppError('Seleccione un nivel educativo válido.')
            if f.type=='grade_ref':
                if not self.db.execute('SELECT 1 FROM grades g JOIN academic_levels l ON l.id=g.level_id WHERE g.name=? AND l.name=? AND g.active=1',(v,str(data.get('nivel','')).strip())).fetchone():raise AppError('El grado no pertenece al nivel seleccionado.')
            if f.type=='section_ref':
                if not self.db.execute('SELECT 1 FROM sections s JOIN grades g ON g.id=s.grade_id JOIN academic_levels l ON l.id=g.level_id WHERE s.year_id=? AND s.name=? AND g.name=? AND l.name=? AND s.active=1',(year,v,str(data.get('grado','')).strip(),str(data.get('nivel','')).strip())).fetchone():raise AppError('La sección no pertenece al grado seleccionado en este año.')
            if f.type=='date':
                try: d=date.fromisoformat(v)
                except ValueError: raise AppError(f'{f.label}: use AAAA-MM-DD.')
                if f.key=='fecha' and not (y['start_date']<=v<=y['end_date']): raise AppError('La fecha debe estar dentro del año escolar seleccionado.')
                if f.key=='nacimiento' and v>y['end_date']: raise AppError('Fecha de nacimiento no válida.')
            if f.type=='time' and not re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d',v): raise AppError(f'{f.label}: use HH:MM.')
            if f.type=='integer':
                if not v.isascii() or not v.isdigit() or int(v)>1_000_000: raise AppError(f'{f.label}: use un entero de 0 a 1000000.')
                v=int(v)
            if f.type=='money':
                try:
                    amount=Decimal(v)
                    if not amount.is_finite() or amount<0 or amount>Decimal('999999999') or amount.as_tuple().exponent < -2: raise InvalidOperation()
                    v=format(amount.quantize(Decimal('.01')),'f')
                except (InvalidOperation,ValueError): raise AppError(f'{f.label}: use un monto positivo con hasta dos decimales, sin comas.')
            if f.type in ('staff_ref','menu_ref'):
                target='staff' if f.type=='staff_ref' else 'menus'
                try: v=int(v)
                except ValueError: raise AppError(f'Seleccione: {f.label}.')
                ref=self.db.execute('SELECT data FROM records WHERE id=? AND year_id=? AND kind=?',(v,year,target)).fetchone()
                if not ref: raise AppError('La referencia no pertenece a este año escolar.')
                rd=json.loads(ref['data'])
                if target=='menus' and rd['servicio']!=data.get('servicio'): raise AppError('El menú debe corresponder al servicio recibido.')
            out[f.key]=v
        if kind=='students' and out['nacimiento']>out['fecha']: raise AppError('El nacimiento debe ser anterior a la matrícula.')
        if kind=='attendance' and out['entrada'] and out['salida'] and out['salida']<out['entrada']: raise AppError('La salida no puede ser anterior a la entrada.')
        if kind=='visits' and out['hora_salida'] and out['hora_salida']<out['hora_entrada']: raise AppError('La hora de salida no puede ser anterior a la entrada.')
        if kind=='receipts' and (out['cantidad']<=0 or out['rechazadas']>out['cantidad']): raise AppError('La cantidad debe ser mayor a cero y cubrir las unidades rechazadas.')
        if kind=='expenses' and Decimal(out['monto'])<=0: raise AppError('El gasto debe ser mayor a cero.')
        return out
    def save_record(self,kind,year,data,rid=None,version=None):
        self.require(kind,write=True); self._open(year); clean=self._validate(kind,year,data)
        try:
            with self.db:
                if rid:
                    old=self.record(kind,year,rid)
                    if version!=old['version']: raise AppError('El registro cambió; cierre y vuelva a abrir el formulario.')
                    # Menús referenciados conservan servicio e identidad; detalles pueden corregirse auditadamente.
                    if kind=='menus' and clean['servicio']!=old['data']['servicio'] and self.db.execute("SELECT 1 FROM records WHERE year_id=? AND kind='receipts' AND json_extract(data,'$.menu')=?",(year,rid)).fetchone(): raise AppError('El servicio de un menú utilizado no se puede cambiar.')
                    cur=self.db.execute('UPDATE records SET data=?,updated_at=?,version=version+1 WHERE id=? AND version=?',(json.dumps(clean,ensure_ascii=False),now(),rid,version))
                    if cur.rowcount!=1: raise AppError('El registro fue modificado por otra sesión.')
                    self._audit('editar_'+kind,year,json.dumps({'id':rid,'antes':old['data'],'despues':clean},ensure_ascii=False))
                else:
                    rid=self.db.execute('INSERT INTO records(year_id,kind,data,created_at,updated_at) VALUES(?,?,?,?,?)',(year,kind,json.dumps(clean,ensure_ascii=False),now(),now())).lastrowid
                    self._audit('crear_'+kind,year,f'id={rid}')
            return rid
        except sqlite3.IntegrityError as e: raise AppError('No se guardó: código duplicado, asistencia ya registrada o año bloqueado.') from e
    def delete_menu(self,year,rid,reason):
        self.require('menus',write=True);self._open(year);reason=' '.join(str(reason).split())
        if len(reason)<5:raise AppError('Escriba un motivo de al menos 5 caracteres.')
        old=self.record('menus',year,rid)
        if old['data'].get('_deleted'):raise AppError('Este menú ya fue eliminado.')
        data=dict(old['data']);data.update({'estado':'Inactivo','_deleted':1,'_deleted_at':now(),'_deleted_by':self.user['username'],'_delete_reason':reason})
        with self.db:
            self.db.execute('UPDATE records SET data=?,updated_at=?,version=version+1 WHERE id=? AND year_id=? AND kind=\'menus\'',(json.dumps(data,ensure_ascii=False),now(),rid,year))
            self._audit('eliminar_menu',year,json.dumps({'id':rid,'nombre':old['data'].get('nombre',''),'motivo':reason},ensure_ascii=False))
    def copy_people(self,source,target,kind):
        self.require(admin=True)
        if kind not in ('students','staff'): raise AppError('Solo se trasladan matrículas o personal.')
        if source==target: raise AppError('Elija años diferentes.')
        self._open(target); target_year=self.year(target); source_year=self.year(source)
        if target_year['label']<=source_year['label']: raise AppError('El año de destino debe ser posterior.')
        items=self.list_records(kind,source)
        codes={r['data']['codigo'] for r in self.list_records(kind,target)}; count=0
        # Una sola transacción; no reutiliza IDs, asistencias ni movimientos.
        with self.db:
            for r in items:
                d=dict(r['data'])
                if d['estado']!='Activo' or d['codigo'] in codes: continue
                if kind=='students': d['fecha']=target_year['start_date']
                d=self._validate(kind,target,d)
                self.db.execute('INSERT INTO records(year_id,kind,data,created_at,updated_at) VALUES(?,?,?,?,?)',(target,kind,json.dumps(d,ensure_ascii=False),now(),now())); count+=1
            self._audit('copiar_'+kind,target,f'origen={source}; cantidad={count}; grado sin promoción automática')
        return count
    def student_profile(self,year,data):
        on=date.today().isoformat(); years,months=age_parts(data['nacimiento'],on);start=self.year(year)['start_date'];start_years,start_months=age_parts(data['nacimiento'],start)
        grade=self.db.execute('SELECT g.min_age,g.max_age FROM grades g JOIN academic_levels l ON l.id=g.level_id WHERE g.name=? AND l.name=?',(data.get('grado',''),data.get('nivel',''))).fetchone()
        expected='Sin rango configurado';matches=None
        if grade and grade['min_age'] is not None and grade['max_age'] is not None:
            expected=f"{grade['min_age']}–{grade['max_age']} años" if grade['min_age']!=grade['max_age'] else f"{grade['min_age']} años"
            matches=grade['min_age']<=start_years<=grade['max_age']
        return {'edad':f'{years} años, {months} meses','edad_inicio_año':f'{start_years} años, {start_months} meses','rango_esperado':expected,'edad_en_rango':matches,'calculada_el':on}
    def _next_code(self,kind,year,offset=0):
        prefix='EST' if kind=='students' else 'PER';label=self.year(year)['label'][:4]
        existing={r['data'].get('codigo','') for r in self.list_records(kind,year)}
        n=1+offset
        while f'{prefix}-{label}-{n:04d}' in existing:n+=1
        return f'{prefix}-{label}-{n:04d}'
    @staticmethod
    def _sheet_headers(kind):
        if kind=='students':return ['Código','Nombre completo','Fecha de nacimiento','Sexo','Documento','Nacionalidad','Dirección','Tutor','Parentesco','Teléfono','Emergencia','Salud','Apoyos','Nivel','Grado','Sección','Fecha de matrícula','Estado','Observaciones']
        if kind=='staff':return ['Código','ID del ponchador','Nombre completo','Documento','Cargo','Teléfono','Dirección','Fecha de ingreso','Vínculo laboral','Salario','Estado','Observaciones']
        raise AppError('La importación está disponible para estudiantes y personal.')
    def create_import_template(self,kind,path):
        self.require(kind);from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill
        wb=Workbook();ws=wb.active;ws.title='Datos';headers=self._sheet_headers(kind);ws.append(headers)
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='176C77')
        ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{chr(64+min(len(headers),26))}1'
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(32,max(14,len(str(col[0].value))+3))
        wb.save(path)
    def import_records(self,kind,year,path,defaults=None):
        self.require(kind,write=True);self._open(year);defaults=dict(defaults or {})
        if kind not in ('students','staff'):raise AppError('Solo se importan estudiantes o personal.')
        p=Path(path)
        if p.suffix.lower()=='.csv':
            with p.open(encoding='utf-8-sig',newline='') as fh:raw=list(csv.reader(fh))
        elif p.suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            ws=load_workbook(p,read_only=True,data_only=True).active;raw=[[c for c in row] for row in ws.iter_rows(values_only=True)]
        else:raise AppError('Seleccione un archivo Excel .xlsx o CSV.')
        if len(raw)<2:raise AppError('El archivo no contiene filas de datos.')
        aliases={
          'codigo':'codigo','id estudiante':'codigo','matricula':'codigo','nombre':'nombre','nombre completo':'nombre','nombres y apellidos':'nombre',
          'fecha de nacimiento':'nacimiento','nacimiento':'nacimiento','sexo':'sexo','genero':'sexo','documento':'documento','cedula':'documento',
          'nacionalidad':'nacionalidad','direccion':'direccion','tutor':'tutor','padre madre o tutor':'tutor','parentesco':'parentesco','telefono':'telefono','telefono del tutor':'telefono',
          'emergencia':'emergencia','salud':'salud','apoyos':'apoyos','nivel':'nivel','nivel educativo':'nivel','grado':'grado','seccion':'seccion',
          'fecha de matricula':'fecha','fecha':'fecha','estado':'estado','observaciones':'notas','notas':'notas','cargo':'cargo','fecha de ingreso':'ingreso','ingreso':'ingreso',
          'vinculo laboral':'vinculo','vinculo':'vinculo','salario':'salario','id del ponchador':'reloj_id','id ponchador':'reloj_id','id reloj':'reloj_id','user id':'reloj_id','id usuario':'reloj_id'}
        mapped=[aliases.get(self._norm(h)) for h in raw[0]]
        if 'nombre' not in mapped:raise AppError('No se encontró la columna Nombre completo.')
        prepared=[];errors=[]
        for index,row in enumerate(raw[1:],2):
            if not any(v not in (None,'') for v in row):continue
            data={k:(v.date().isoformat() if isinstance(v,datetime) else v.isoformat() if isinstance(v,date) else str(v).strip()) for k,v in zip(mapped,row) if k and v not in (None,'')};data={**defaults,**data}
            if not data.get('codigo'):data['codigo']=self._next_code(kind,year,len(prepared))
            if kind=='students':
                level=self._match_level(data.get('nivel','')) if data.get('nivel') else None
                if level:data['nivel']=level['name']
                grade=self._match_grade(data.get('grado',''),data.get('nivel',''))
                if grade:data['grado']=grade['name'];data['nivel']=grade['level']
                data.setdefault('sexo','No especificado');data.setdefault('tutor','Pendiente de completar');data.setdefault('telefono','Pendiente de completar')
                data.setdefault('seccion','A');data['seccion']=data['seccion'].upper();data.setdefault('fecha',self.year(year)['start_date']);data.setdefault('estado','Activo')
            else:
                data.setdefault('ingreso',self.year(year)['start_date']);data.setdefault('vinculo','Temporal');data.setdefault('salario','0');data.setdefault('estado','Activo')
            try:prepared.append(self._validate(kind,year,data))
            except AppError as e:errors.append(f'Fila {index}: {e}')
        if errors:raise AppError('No se importó ninguna fila. Corrige lo siguiente:\n'+'\n'.join(errors[:12])+('\n…' if len(errors)>12 else ''))
        if not prepared:raise AppError('No se encontraron datos para importar.')
        try:
            with self.db:
                for data in prepared:self.db.execute('INSERT INTO records(year_id,kind,data,created_at,updated_at) VALUES(?,?,?,?,?)',(year,kind,json.dumps(data,ensure_ascii=False),now(),now()))
                self._audit('importar_'+kind,year,f'archivo={p.name}; cantidad={len(prepared)}')
        except sqlite3.IntegrityError as e:raise AppError('La importación contiene códigos duplicados o datos repetidos.') from e
        return len(prepared)
    def _prepare_attendance_import(self,year,path):
        self.require('attendance');y=self.year(year)
        try:parsed=parse_punch_file(path)
        except PunchFormatError as e:raise AppError(str(e)) from e
        staff=self.list_records('staff',year)
        indexes={key:{} for key in ('reloj_id','codigo','documento','nombre')}
        for row in staff:
            for key in indexes:
                value=self._norm(row['data'].get(key,''))
                if value:indexes[key].setdefault(value,[]).append(row)

        def resolve(session):
            attempts=(
                ('reloj_id',session.get('clock_id','')),('codigo',session.get('code','')),
                ('documento',session.get('document','')),('nombre',session.get('name','')),
                ('codigo',session.get('clock_id','')),('reloj_id',session.get('code','')),
            )
            for field,value in attempts:
                matches=indexes[field].get(self._norm(value),[]) if value else []
                if len(matches)==1:return matches[0]
                if len(matches)>1:return None
            return None

        settings=self.settings();expected=settings.get('attendance_start','08:00');grace=int(settings.get('attendance_grace','10') or 0)
        expected_minutes=int(expected[:2])*60+int(expected[3:])+grace
        ready=[];unmatched=[];errors=list(parsed['errors'])
        for session in parsed['sessions']:
            if not (y['start_date']<=session['date']<=y['end_date']):
                errors.append(f"Filas {','.join(map(str,session['source_rows']))}: fecha {session['date']} fuera del año escolar.")
                continue
            person=resolve(session)
            if not person:
                identity=session.get('clock_id') or session.get('code') or session.get('document') or session.get('name') or 'sin identificar'
                unmatched.append(f"{identity} · {session['date']}")
                continue
            if session['entry'] and session['exit'] and session['exit']<session['entry']:
                errors.append(f"{person['data']['nombre']} · {session['date']}: salida anterior a la entrada.")
                continue
            late=False
            if session['entry']:
                late=int(session['entry'][:2])*60+int(session['entry'][3:])>expected_minutes
            data={'personal':person['id'],'fecha':session['date'],'estado':'Tardanza' if late else 'Presente',
                  'entrada':session['entry'],'salida':session['exit'],
                  'notas':f"Importado de {Path(path).name}; {session['mark_count']} marcación(es); filas {','.join(map(str,session['source_rows']))}."}
            ready.append(self._validate('attendance',year,data))
        return {**parsed,'ready':ready,'unmatched':unmatched,'errors':errors,'matched_sessions':len(ready)}
    def analyze_attendance_file(self,year,path):
        result=self._prepare_attendance_import(year,path)
        return {key:result[key] for key in ('header_row','recognized_columns','source_rows','source_marks','matched_sessions','unmatched','errors')}
    def import_attendance_file(self,year,path):
        self.require('attendance',write=True);self._open(year);result=self._prepare_attendance_import(year,path)
        if not result['ready']:
            raise AppError('No hay jornadas que coincidan con el personal de este año. Revise el ID del ponchador o el código del personal.')
        existing={(row['data']['personal'],row['data']['fecha']):row for row in self.list_records('attendance',year)}
        created=updated=unchanged=0
        with self.db:
            for data in result['ready']:
                old=existing.get((data['personal'],data['fecha']))
                if not old:
                    rid=self.db.execute('INSERT INTO records(year_id,kind,data,created_at,updated_at) VALUES(?,?,?,?,?)',(year,'attendance',json.dumps(data,ensure_ascii=False),now(),now())).lastrowid
                    self._audit('importar_asistencia_ponchador',year,f'id={rid}; personal={data["personal"]}; fecha={data["fecha"]}')
                    created+=1;continue
                merged=dict(old['data']);entries=[value for value in (merged.get('entrada',''),data.get('entrada','')) if value];exits=[value for value in (merged.get('salida',''),data.get('salida','')) if value]
                merged['entrada']=min(entries) if entries else '';merged['salida']=max(exits) if exits else ''
                if merged.get('estado') in ('Presente','Tardanza'):
                    settings=self.settings();expected=settings.get('attendance_start','08:00');grace=int(settings.get('attendance_grace','10') or 0);limit=int(expected[:2])*60+int(expected[3:])+grace
                    entry_minutes=int(merged['entrada'][:2])*60+int(merged['entrada'][3:]) if merged['entrada'] else 0
                    merged['estado']='Tardanza' if merged['entrada'] and entry_minutes>limit else 'Presente'
                if data['notas'] not in merged.get('notas',''):merged['notas']=(merged.get('notas','')+'\n'+data['notas']).strip()
                merged=self._validate('attendance',year,merged)
                if merged==old['data']:unchanged+=1;continue
                self.db.execute('UPDATE records SET data=?,updated_at=?,version=version+1 WHERE id=? AND version=?',(json.dumps(merged,ensure_ascii=False),now(),old['id'],old['version']))
                self._audit('actualizar_asistencia_ponchador',year,f'id={old["id"]}; personal={data["personal"]}; fecha={data["fecha"]}');updated+=1
            self._audit('resumen_importacion_ponchador',year,f'archivo={Path(path).name}; creados={created}; actualizados={updated}; sin_cambios={unchanged}; no_coinciden={len(result["unmatched"])}; errores={len(result["errors"])}')
        return {'created':created,'updated':updated,'unchanged':unchanged,'matched_sessions':result['matched_sessions'],'source_marks':result['source_marks'],'unmatched':result['unmatched'],'errors':result['errors']}
    @staticmethod
    def _attachment_payload(path):
        p=Path(path)
        if p.suffix.lower() not in {'.pdf','.jpg','.jpeg','.png','.xlsx','.docx','.txt'}: raise AppError('Adjunte PDF, imagen, Excel, Word o TXT.')
        if p.stat().st_size>25*1024*1024: raise AppError('El adjunto debe ser menor a 25 MB.')
        content=p.read_bytes()
        if len(content)>25*1024*1024 or not content: raise AppError('Adjunto vacío o demasiado grande.')
        mime='application/pdf' if content.startswith(b'%PDF-') else 'image/png' if content.startswith(b'\x89PNG\r\n\x1a\n') else 'image/jpeg' if content.startswith(b'\xff\xd8\xff') else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if p.suffix.lower()=='.xlsx' and content.startswith(b'PK') else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' if p.suffix.lower()=='.docx' and content.startswith(b'PK') else 'text/plain' if p.suffix.lower()=='.txt' else None
        if not mime: raise AppError('El contenido del archivo no coincide con su formato.')
        return p.name,mime,hashlib.sha256(content).hexdigest(),content
    def attach_many(self,kind,year,rid,paths):
        self.require(kind,write=True);self._open(year);self.record(kind,year,rid);paths=list(paths)
        if not paths:raise AppError('Seleccione uno o varios archivos.')
        if len(paths)>20:raise AppError('Puede agregar hasta 20 archivos en una sola selección.')
        prepared=[self._attachment_payload(path) for path in paths];ids=[]
        with self.db:
            for name,mime,digest,content in prepared:
                aid=self.db.execute('INSERT INTO attachments(record_id,year_id,name,mime,sha256,content) VALUES(?,?,?,?,?,?)',(rid,year,name,mime,digest,content)).lastrowid
                ids.append(aid);self._audit('adjuntar_evidencia',year,f'registro={rid}; adjunto={aid}')
            self._audit('adjuntar_evidencias_lote',year,f'registro={rid}; cantidad={len(ids)}')
        return ids
    def attach(self,kind,year,rid,path):
        return self.attach_many(kind,year,rid,[path])[0]
    def attachments(self,kind,year,rid):
        self.record(kind,year,rid)
        return [dict(r) for r in self.db.execute('SELECT id,name,mime,length(content) AS size FROM attachments WHERE record_id=? AND year_id=?',(rid,year))]
    def attachment_bytes(self,kind,year,rid,aid,audit=True):
        self.record(kind,year,rid)
        r=self.db.execute('SELECT content,sha256 FROM attachments WHERE id=? AND record_id=? AND year_id=?',(aid,rid,year)).fetchone()
        if not r or hashlib.sha256(r['content']).hexdigest()!=r['sha256']: raise AppError('Adjunto no encontrado o dañado.')
        if audit:
            with self.db:self._audit('exportar_adjunto',year,f'adjunto={aid}')
        return r['content']
    def dashboard(self,year):
        self.require(); self.year(year); result={}
        role=self.user['role']
        if 'students' in ACCESS[role]:
            result['Estudiantes activos']=sum(r['data']['estado']=='Activo' for r in self.list_records('students',year))
            result['Secciones activas']=len(self.sections(year))
        if 'staff' in ACCESS[role]: result['Personal activo']=sum(r['data']['estado']=='Activo' for r in self.list_records('staff',year))
        if 'attendance' in ACCESS[role]:
            rows=self.list_records('attendance',year); today=date.today().isoformat()
            result['Presentes hoy']=sum(r['data']['fecha']==today and r['data']['estado'] in ('Presente','Tardanza') for r in rows)
        if 'receipts' in ACCESS[role]:
            rows=[r['data'] for r in self.list_records('receipts',year) if r['data']['estado']!='Anulado']
            configured=[item['name'] for item in self.food_services()];services=configured+sorted({r['servicio'] for r in rows}-set(configured))
            for service in services: result[f'{service} · unidades aceptadas']=sum(r['cantidad']-r['rechazadas'] for r in rows if r['servicio']==service)
            result['Recepciones con incidencias']=sum(bool(r.get('incidencias')) or bool(r.get('observaciones')) or r['rechazadas']>0 for r in rows)
        if 'expenses' in ACCESS[role]:
            total=sum((Decimal(r['data']['monto']) for r in self.list_records('expenses',year) if r['data']['estado']!='Anulado'),Decimal(0))
            result['Gastos RD$']=f'{total:,.2f}'
        if 'visits' in ACCESS[role]: result['Visitas de hoy']=sum(r['data']['fecha']==date.today().isoformat() for r in self.list_records('visits',year))
        if 'facilities' in ACCESS[role]: result['Planta física · pendientes']=sum(r['data']['estado']!='Resuelto' for r in self.list_records('facilities',year))
        return result
    def audit(self,year=None):
        self.require()
        if self.user['role'] not in ('admin','direccion'): raise AppError('Auditoría reservada a Dirección y Administración.')
        sql='SELECT a.*,COALESCE(u.username,\'sistema\') AS usuario FROM audit a LEFT JOIN users u ON u.id=a.user_id'
        args=()
        if year is not None: sql+=' WHERE a.year_id=?'; args=(year,)
        rows=[dict(x) for x in self.db.execute(sql+' ORDER BY a.id DESC',args)]
        for row in rows:
            try:row['at_local']=datetime.fromisoformat(row['at']).astimezone().strftime('%d/%m/%Y %I:%M:%S %p')
            except ValueError:row['at_local']=row['at']
        return rows
    def log_export(self,kind,year):
        self.require(kind)
        with self.db: self._audit('exportar_'+kind,year)
