import json,sqlite3,time
from decimal import Decimal
import pytest
from aulalocal.core import Store,AppError
from conftest import PASSWORD,student,staff,menu,receipt,expense

def test_persistence_and_school_config(store):
    rid=store.save_record('students',1,student())
    store.save_settings({'centro':'Centro editable','director':'Director editable'})
    other=Store(store.root)
    try:
        other.login('director',PASSWORD)
        assert other.record('students',1,rid)['data']['nombre']=='María de Prueba'
        assert other.settings()['centro']=='Centro editable'
    finally:other.close()

def test_year_isolation_and_copies(store):
    rid=store.save_record('students',1,student());store.save_record('expenses',1,expense())
    y=store.create_year('2027-2028','2027-08-01','2028-07-31')
    assert store.list_records('students',y)==[]
    with pytest.raises(AppError):store.record('students',y,rid)
    assert store.copy_people(1,y,'students')==1
    assert store.copy_people(1,y,'students')==0
    assert store.list_records('expenses',y)==[]
    r=store.list_records('students',y)[0];d=r['data'];d['grado']='2.º de Secundaria'
    store.save_record('students',y,d,r['id'],r['version'])
    assert store.record('students',1,rid)['data']['grado']=='1.º de Secundaria'

def test_freeze_db_and_service_and_reopen(store):
    rid=store.save_record('students',1,student())
    store.transition(1,'congelado','Cierre del año escolar',PASSWORD)
    with pytest.raises(AppError):store.save_record('students',1,student('E002'))
    with pytest.raises(sqlite3.IntegrityError):
        with store.db:store.db.execute('UPDATE records SET updated_at=? WHERE id=?',('x',rid))
    with pytest.raises(AppError):store.transition(1,'abierto','Corrección autorizada','incorrecta')
    store.transition(1,'archivado','Archivo del año escolar',PASSWORD)
    with pytest.raises(AppError):store.transition(1,'abierto','Reapertura autorizada',PASSWORD)
    store.transition(1,'congelado','Revisión del archivo histórico',PASSWORD)
    store.transition(1,'abierto','Reapertura autorizada',PASSWORD)
    store.save_record('students',1,student('E002'))
    assert len(store.list_records('students',1))==2

def test_duplicate_and_optimistic_concurrency(store):
    rid=store.save_record('students',1,student())
    with pytest.raises(AppError):store.save_record('students',1,student())
    r=store.record('students',1,rid);d=r['data'];d['nombre']='Nombre corregido'
    store.save_record('students',1,d,rid,r['version'])
    with pytest.raises(AppError):store.save_record('students',1,d,rid,r['version'])
    assert any('María de Prueba' in a['detail'] for a in store.audit())

@pytest.mark.parametrize('role,allowed,denied',[('alimentacion','menus','students'),('contabilidad','expenses','staff'),('secretaria','students','expenses')])
def test_roles(store,role,allowed,denied):
    store.save_user(role,'Usuario prueba',role,PASSWORD);store.login(role,PASSWORD)
    assert store.list_records(allowed,1)==[]
    with pytest.raises(AppError):store.list_records(denied,1)
    with pytest.raises(AppError):store.create_year('2027-2028','2027-08-01','2028-07-31')

def test_read_only_and_persistent_session(store):
    store.save_user('consulta','Consulta','consulta',PASSWORD);store.login('consulta',PASSWORD)
    with pytest.raises(AppError):store.save_record('students',1,student())
    store.last_active=time.monotonic()-901
    assert store.years()
    other=Store(store.root)
    try:assert other.user['username']=='consulta'
    finally:other.close()

def test_lockout_and_no_plaintext_password(store):
    for _ in range(5):
        with pytest.raises(AppError):store.login('director','incorrecta')
    with pytest.raises(AppError,match='bloqueada'):store.login('director',PASSWORD)
    row=store.db.execute('SELECT password_hash,salt FROM users').fetchone()
    assert len(row['password_hash'])==32 and len(row['salt'])==16
    assert PASSWORD.encode() not in store.path.read_bytes()

def test_attendance_reference_and_duplicate(store):
    pid=store.save_record('staff',1,staff());a=dict(personal=pid,fecha='2026-09-01',estado='Presente',entrada='08:00',salida='16:00')
    store.save_record('attendance',1,a)
    with pytest.raises(AppError):store.save_record('attendance',1,a)
    y=store.create_year('2027-2028','2027-08-01','2028-07-31');a['fecha']='2027-09-01'
    with pytest.raises(AppError):store.save_record('attendance',y,a)

def test_food_menu_quantities_and_dashboard(store):
    mid=store.save_record('menus',1,menu());rid=store.save_record('receipts',1,receipt(mid))
    assert store.dashboard(1)['Almuerzo · unidades aceptadas']==118
    bad=receipt(mid);bad['rechazadas']='121'
    with pytest.raises(AppError):store.save_record('receipts',1,bad)
    bad=receipt(mid);bad['servicio']='Leche'
    with pytest.raises(AppError):store.save_record('receipts',1,bad)
    r=store.record('receipts',1,rid);r['data']['estado']='Anulado';store.save_record('receipts',1,r['data'],rid,r['version'])
    assert store.dashboard(1)['Almuerzo · unidades aceptadas']==0

def test_configurable_food_types_and_controlled_menu_deletion(store):
    store.save_food_service('Merienda escolar')
    assert 'Merienda escolar' in [item['name'] for item in store.food_services()]
    custom=menu('Merienda escolar');custom['nombre']='Fruta y yogur';mid=store.save_record('menus',1,custom)
    received=receipt(mid);received.update(servicio='Merienda escolar',observaciones='Empaque revisado');rid=store.save_record('receipts',1,received)
    assert store.dashboard(1)['Merienda escolar · unidades aceptadas']==118
    store.delete_menu(1,mid,'Menú sustituido por el nuevo ciclo')
    assert store.list_records('menus',1)==[]
    assert store.list_records('menus',1,include_deleted=True)[0]['data']['_deleted']==1
    assert store.record('receipts',1,rid)['data']['menu']==mid
    assert any(item['action']=='eliminar_menu' for item in store.audit())

def test_exact_money_and_no_invalid_amount(store):
    store.save_record('expenses',1,expense('0.10'));store.save_record('expenses',1,expense('0.20'))
    assert store.dashboard(1)['Gastos RD$']=='0.30'
    for amount in ['-1','NaN','Infinity','12.345','1,000','0']:
        with pytest.raises(AppError):store.save_record('expenses',1,expense(amount))

def test_dates_enforced(store):
    for d in ['2026-07-31','2027-08-01','2026-99-01']:
        s=student();s['fecha']=d
        with pytest.raises(AppError):store.save_record('students',1,s)

def test_attachments_and_frozen_files(store,tmp_path):
    from PIL import Image
    p=tmp_path/'evidencia.png';Image.new('RGB',(8,8),'blue').save(p)
    rid=store.save_record('students',1,student());aid=store.attach('students',1,rid,p)
    assert store.attachment_bytes('students',1,rid,aid)==p.read_bytes()
    bad=tmp_path/'falso.pdf';bad.write_text('no PDF')
    with pytest.raises(AppError):store.attach('students',1,rid,bad)
    store.transition(1,'congelado','Cierre de prueba formal',PASSWORD)
    with pytest.raises(AppError):store.attach('students',1,rid,p)
    assert store.attachments('students',1,rid)[0]['name']=='evidencia.png'

def test_audit_cannot_be_deleted(store):
    with pytest.raises(sqlite3.IntegrityError):
        with store.db:store.db.execute('DELETE FROM audit')

def test_last_admin_cannot_disable_self(store):
    with pytest.raises(AppError):store.save_user('director','Director','consulta',user_id=store.user['id'])

def test_academic_catalog_age_and_sections(store):
    levels=store.academic_levels();assert [x['name'] for x in levels][:2]==['Nivel Inicial','Nivel Primario']
    kinder=next(x for x in store.grades() if x['name']=='Kínder')
    assert kinder['min_age']==4 and kinder['max_age']==4
    section=next(x for x in store.sections(1,kinder['id']) if x['name']=='A')
    s=student();s.update(nacimiento='2022-07-01',nivel='Nivel Inicial',grado='Kínder',seccion=section['name'])
    store.save_record('students',1,s);profile=store.student_profile(1,s)
    assert 'años' in profile['edad'] and profile['rango_esperado']=='4 años'

def test_excel_bulk_import_students_and_staff(store,tmp_path):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.append(['Nombre completo','Fecha de nacimiento']);ws.append(['Luis Uno','2022-07-01']);ws.append(['Ana Dos','2022-06-02'])
    path=tmp_path/'estudiantes.xlsx';wb.save(path)
    count=store.import_records('students',1,path,{'nivel':'Nivel Inicial','grado':'Kínder','seccion':'A'})
    assert count==2 and {r['data']['grado'] for r in store.list_records('students',1)}=={'Kínder'}
    template=tmp_path/'personal.xlsx';store.create_import_template('staff',template);assert template.read_bytes().startswith(b'PK')

def test_punch_clock_import_groups_marks_matches_staff_and_is_idempotent(store,tmp_path):
    from openpyxl import Workbook
    first=staff();first['reloj_id']='42';pid=store.save_record('staff',1,first)
    second={**staff(),'codigo':'P002','reloj_id':'43','nombre':'Berta de Prueba'};pid2=store.save_record('staff',1,second)
    wb=Workbook();ws=wb.active;ws.append(['Reporte biométrico']);ws.append(['ID ponchador','Nombre','Fecha y hora']);ws.append(['42','Ana de Prueba','01/09/2026 08:12']);ws.append(['42','Ana de Prueba','01/09/2026 16:07']);ws.append(['999','Sin coincidencia','01/09/2026 07:55']);path=tmp_path/'marcaciones.xlsx';wb.save(path)
    analysis=store.analyze_attendance_file(1,path);assert analysis['header_row']==2 and analysis['source_marks']==3;assert analysis['matched_sessions']==1 and len(analysis['unmatched'])==1
    imported=store.import_attendance_file(1,path);assert imported['created']==1 and imported['updated']==0
    attendance=store.list_records('attendance',1)[0]['data'];assert attendance['personal']==pid and attendance['entrada']=='08:12' and attendance['salida']=='16:07' and attendance['estado']=='Tardanza'
    again=store.import_attendance_file(1,path);assert again['unchanged']==1 and len(store.list_records('attendance',1))==1
    direct=Workbook();sheet=direct.active;sheet.append(['User ID','Date','Check In','Check Out']);sheet.append(['43','2026-09-02','07:58','16:01']);direct_path=tmp_path/'entrada-salida.xlsx';direct.save(direct_path)
    result=store.import_attendance_file(1,direct_path);row=next(r['data'] for r in store.list_records('attendance',1) if r['data']['personal']==pid2);assert result['created']==1 and row['entrada']=='07:58' and row['salida']=='16:01' and row['estado']=='Presente'

def test_punch_clock_semicolon_csv_and_corrupt_file(store,tmp_path):
    person=staff();person['reloj_id']='77';store.save_record('staff',1,person)
    csv_path=tmp_path/'reloj.csv';csv_path.write_text('ID ponchador;Fecha;Hora\n77;03/09/2026;07:59\n77;03/09/2026;16:02\n',encoding='utf-8')
    result=store.import_attendance_file(1,csv_path);assert result['created']==1;row=store.list_records('attendance',1)[0]['data'];assert row['entrada']=='07:59' and row['salida']=='16:02'
    corrupt=tmp_path/'dañado.xlsx';corrupt.write_text('esto no es excel',encoding='utf-8')
    with pytest.raises(AppError,match='dañado'):
        store.analyze_attendance_file(1,corrupt)

def test_multiple_food_photos_are_atomic_and_audited(store,tmp_path):
    from PIL import Image
    mid=store.save_record('menus',1,menu());rid=store.save_record('receipts',1,receipt(mid));first=tmp_path/'almuerzo-frente.jpg';second=tmp_path/'almuerzo-bandejas.png';Image.new('RGB',(640,480),'orange').save(first);Image.new('RGB',(800,600),'green').save(second)
    ids=store.attach_many('receipts',1,rid,[first,second]);assert len(ids)==2 and len(store.attachments('receipts',1,rid))==2
    bad=tmp_path/'imagen-falsa.jpg';bad.write_text('no es una imagen',encoding='utf-8')
    with pytest.raises(AppError):store.attach_many('receipts',1,rid,[first,bad])
    assert len(store.attachments('receipts',1,rid))==2 and any(item['action']=='adjuntar_evidencias_lote' for item in store.audit())

def test_recovery_key_and_explicit_logout(store):
    code=store.create_recovery_key(PASSWORD);store.logout();assert not store.session_file.exists()
    new_code=store.recover_access(code,'direccion2','Nueva-clave-segura-2026');store.login('direccion2','Nueva-clave-segura-2026')
    assert store.user['username']=='direccion2'
    assert new_code.startswith('AULA-')
    store.logout();store.recover_access(new_code,'direccion3','Otra-clave-segura-2026')

def test_facilities_visits_and_office_attachments(store,tmp_path):
    visit=dict(fecha='2026-09-01',hora_entrada='09:00',nombre='Visitante',motivo='Supervisión',recibido_por='Dirección')
    assert store.save_record('visits',1,visit)
    facility=dict(fecha='2026-09-01',area='Aula 1',condicion='Regular',hallazgo='Pintura deteriorada',estado='Pendiente')
    rid=store.save_record('facilities',1,facility)
    from openpyxl import Workbook
    path=tmp_path/'evidencia.xlsx';Workbook().save(path);assert store.attach('facilities',1,rid,path)

def test_v1_upgrade_preserves_and_normalizes_students(tmp_path):
    root=tmp_path/'version-anterior';old=Store(root);old.bootstrap('director',PASSWORD,'Directora','Centro')
    year=old.create_year('2026-2027','2026-08-01','2027-07-31');rid=old.save_record('students',year,student())
    evidence=tmp_path/'nota.txt';evidence.write_text('Documento conservado',encoding='utf-8');aid=old.attach('students',year,rid,evidence);old.close()
    legacy=student();legacy.pop('nivel');legacy['grado']='1.º secundaria'
    with sqlite3.connect(root/'aulalocal.sqlite3') as db:
        db.execute("UPDATE meta SET value='1' WHERE key='schema_version'")
        db.execute('UPDATE records SET data=? WHERE id=?',(json.dumps(legacy,ensure_ascii=False),rid))
    upgraded=Store(root)
    try:
        record=upgraded.record('students',year,rid)
        assert record['data']['nivel']=='Nivel Secundario'
        assert record['data']['grado']=='1.º de Secundaria'
        assert upgraded.attachment_bytes('students',year,rid,aid)==b'Documento conservado'
        assert upgraded.db.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()[0]=='2'
    finally:upgraded.close()
