import io,zipfile,socket
from pathlib import Path
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from aulalocal.core import AppError
from aulalocal.backup import create_backup,restore_backup
from aulalocal import reports
from conftest import PASSWORD,student,staff,expense,menu,receipt

def test_backup_round_trip_with_attachments(store,tmp_path):
    rid=store.save_record('students',1,student());pdf=tmp_path/'expediente.pdf';reports.record_pdf(store,'students',1,rid,pdf);aid=store.attach('students',1,rid,pdf)
    backup=tmp_path/'copia.aulabackup';create_backup(store,backup,PASSWORD)
    assert not backup.read_bytes().startswith(b'SQLite')
    assert b'Mar' not in backup.read_bytes()[:32]
    store.save_record('students',1,student('E002'))
    recovery=restore_backup(store,backup,PASSWORD,PASSWORD)
    assert recovery.exists() and store.user is None
    store.login('director',PASSWORD);assert len(store.list_records('students',1))==1
    assert store.attachment_bytes('students',1,rid,aid)==pdf.read_bytes()

def test_wrong_password_and_tamper_leave_current_data(store,tmp_path):
    store.save_record('students',1,student());backup=tmp_path/'copia.aulabackup';create_backup(store,backup,PASSWORD)
    with pytest.raises(AppError):restore_backup(store,backup,'contraseña-falsa',PASSWORD)
    raw=bytearray(backup.read_bytes());raw[-1]^=1;backup.write_bytes(raw)
    with pytest.raises(AppError):restore_backup(store,backup,PASSWORD,PASSWORD)
    assert len(store.list_records('students',1))==1

def test_backup_permission(store,tmp_path):
    store.save_user('lector','Lector','consulta',PASSWORD);store.login('lector',PASSWORD)
    with pytest.raises(AppError):create_backup(store,tmp_path/'copy',PASSWORD)

def test_pdf_xlsx_and_formula_injection_offline(store,tmp_path,monkeypatch):
    def no_network(*a,**kw):raise AssertionError('La aplicación intentó usar internet')
    monkeypatch.setattr(socket,'socket',no_network)
    rid=store.save_record('students',1,student(name='=HYPERLINK("https://ejemplo.test")'))
    path=tmp_path/'reporte.xlsx';reports.report_xlsx(store,'students',1,path)
    wb=load_workbook(path);assert wb.active['B4'].data_type=='s';assert wb.active['B4'].value.startswith('=HYPERLINK')
    pdf=tmp_path/'reporte.pdf';reports.report_pdf(store,'students',1,pdf)
    assert len(PdfReader(pdf).pages)>=1 and '2026-2027' in PdfReader(pdf).pages[0].extract_text()
    letter=tmp_path/'constancia.pdf';reports.letter_pdf(store,'students',1,rid,letter)
    assert 'CERTIFICACI' in PdfReader(letter).pages[0].extract_text()

def test_excel_money_and_filter(store,tmp_path):
    store.save_record('expenses',1,expense('10.25'));store.save_record('expenses',1,{**expense('99.00'),'concepto':'Otro'})
    path=tmp_path/'gastos.xlsx';reports.report_xlsx(store,'expenses',1,path,'Materiales de prueba')
    ws=load_workbook(path).active
    assert ws.max_row==4;assert ws['F4'].value==10.25;assert ws['F4'].data_type=='n'

def test_student_section_filter_and_attendance_supervision_reports(store,tmp_path):
    first=student();store.save_record('students',1,first);grade=next(item for item in store.grades() if item['name']=='1.º de Secundaria');store.save_section(1,grade['id'],'B');second=student('E002','Estudiante de Sección B');second['seccion']='B';store.save_record('students',1,second)
    filtered=tmp_path/'seccion-b.xlsx';reports.report_xlsx(store,'students',1,filtered,filters={'nivel':'Nivel Secundario','grado':'1.º de Secundaria','seccion':'B'});ws=load_workbook(filtered).active;assert ws.max_row==4 and ws['B4'].value=='Estudiante de Sección B' and 'Sección: B' in ws['A2'].value
    filtered_pdf=tmp_path/'seccion-b.pdf';reports.report_pdf(store,'students',1,filtered_pdf,filters={'grado':'1.º de Secundaria','seccion':'B'});filtered_text=''.join(page.extract_text() for page in PdfReader(filtered_pdf).pages);assert 'Estudiante de' in filtered_text and 'Sección B' in filtered_text and 'E001' not in filtered_text
    person=staff();person['reloj_id']='42';pid=store.save_record('staff',1,person);store.save_record('attendance',1,dict(personal=pid,fecha='2026-09-01',estado='Tardanza',entrada='08:12',salida='16:05',notas='Marcación biométrica'))
    xlsx=tmp_path/'supervision.xlsx';reports.report_attendance_xlsx(store,1,xlsx,'2026-09-01','2026-09-30');workbook=load_workbook(xlsx);assert workbook.sheetnames==['Resumen','Detalle'];assert workbook['Resumen']['C4'].value=='Ana de Prueba' and workbook['Resumen']['G4'].value==1 and workbook['Detalle']['F4'].value=='Martes'
    pdf=tmp_path/'supervision.pdf';reports.report_attendance_pdf(store,1,pdf,'2026-09-01','2026-09-30');text=''.join(page.extract_text() for page in PdfReader(pdf).pages);assert 'Informe de asistencia del personal' in text and 'Resumen por persona' in text and 'Detalle diario' in text

def test_food_reports_with_optional_multiple_photos(store,tmp_path):
    from PIL import Image
    lunch=store.save_record('menus',1,{**menu('Almuerzo'),'nombre':'Arroz y pollo'});bread=store.save_record('menus',1,{**menu('Pan'),'nombre':'Pan escolar'});lunch_receipt=store.save_record('receipts',1,receipt(lunch));bread_data=receipt(bread);bread_data.update(servicio='Pan',cantidad='80',rechazadas='3',hora='08:15');bread_receipt=store.save_record('receipts',1,bread_data)
    photos=[]
    for index,color in enumerate(('orange','green','blue'),1):path=tmp_path/f'evidencia-{index}.jpg';Image.new('RGB',(1200,800),color).save(path,quality=92);photos.append(path)
    store.attach_many('receipts',1,lunch_receipt,photos[:2]);store.attach_many('receipts',1,bread_receipt,photos[2:]);data=reports.food_report_data(store,1,'2026-09-01','2026-09-30');totals={row['servicio']:row for row in data['summary']};assert data['photo_count']==3 and totals['Almuerzo']['aceptadas']==118 and totals['Pan']['aceptadas']==77
    plain=tmp_path/'alimentacion-sin-fotos.pdf';reports.report_food_pdf(store,1,plain,'2026-09-01','2026-09-30',include_photos=False);photographed=tmp_path/'alimentacion-con-fotos.pdf';reports.report_food_pdf(store,1,photographed,'2026-09-01','2026-09-30',include_photos=True);delivery=tmp_path/'fichas-con-fotos.pdf';reports.report_food_pdf(store,1,delivery,'2026-09-01','2026-09-30',include_photos=True,layout='delivery')
    assert photographed.stat().st_size>plain.stat().st_size and len(PdfReader(delivery).pages)>=3;photo_text=''.join(page.extract_text() for page in PdfReader(photographed).pages);assert 'Evidencias fotográficas' in photo_text and 'Arroz y pollo' in photo_text
    xlsx=tmp_path/'alimentacion.xlsx';reports.report_food_xlsx(store,1,xlsx,'2026-09-01','2026-09-30');workbook=load_workbook(xlsx);assert workbook.sheetnames==['Resumen','Recepciones'];evidence_counts={workbook['Recepciones'].cell(row,4).value:workbook['Recepciones'].cell(row,15).value for row in range(4,workbook['Recepciones'].max_row+1)};assert evidence_counts=={'Pan':1,'Almuerzo':2}

def test_frozen_institution_preserved(store,tmp_path):
    store.save_record('students',1,student());store.transition(1,'congelado','Cierre autorizado del año',PASSWORD)
    store.save_settings({'centro':'Nombre nuevo','director':'Director nuevo'})
    assert reports.institution(store,1)['centro']=='Centro de prueba'

def test_year_archive_contains_only_selected_year(store,tmp_path):
    store.save_record('students',1,student());y=store.create_year('2027-2028','2027-08-01','2028-07-31');store.copy_people(1,y,'students')
    z=tmp_path/'año.zip';reports.export_year(store,1,z)
    with zipfile.ZipFile(z) as archive:
        assert 'students.xlsx' in archive.namelist()
        assert '2027-2028' not in archive.read('año.json').decode()
        assert '2027-08-01' not in archive.read('students.json').decode()

def test_long_text_pdf_paginates(store,tmp_path):
    d=expense();d['concepto']='Observación extensa. '*400;store.save_record('expenses',1,d)
    path=tmp_path/'largo.pdf';reports.report_pdf(store,'expenses',1,path)
    assert len(PdfReader(path).pages)>1

def test_export_cannot_overwrite_active_database(store):
    store.save_record('students',1,student())
    with pytest.raises(AppError):reports.report_pdf(store,'students',1,store.path)
    assert len(store.list_records('students',1))==1
